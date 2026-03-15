"""POST /analyze — LLM-powered field detection in documents."""

import logging
from io import BytesIO

from docx import Document
from fastapi import APIRouter, Depends, File, Form, HTTPException, UploadFile

from app.clients.storage_client import get_storage_client
from app.clients.weaviate_client import get_weaviate_client
from app.core.config import get_settings
from app.core.security import verify_api_key
from app.schemas.analyze import AnalyzeResponse
from app.services.analyzer import get_analyzer
from app.services.pdf_extractor import extract_text_from_pdf, extract_widgets_from_pdf, has_form_fields
from app.services.session_store import get_session_store
from app.schemas.session import SessionStatus

logger = logging.getLogger(__name__)
router = APIRouter()


def _extract_text_from_docx(docx_bytes: bytes) -> str:
    """Extract plain text from DOCX bytes."""
    doc = Document(BytesIO(docx_bytes))
    parts = []

    for para in doc.paragraphs:
        if para.text.strip():
            parts.append(para.text)

    for table in doc.tables:
        for row in table.rows:
            row_texts = [cell.text.strip() for cell in row.cells if cell.text.strip()]
            if row_texts:
                parts.append(" | ".join(row_texts))

    return "\n".join(parts)


def _extract_title_from_docx(docx_bytes: bytes, fallback: str = "document") -> str:
    """Extract document title from DOCX core properties or first heading."""
    try:
        doc = Document(BytesIO(docx_bytes))
        # Try core properties
        if doc.core_properties.title:
            return doc.core_properties.title

        # Try first heading
        for para in doc.paragraphs:
            if para.style and para.style.name and "Heading" in para.style.name:
                if para.text.strip():
                    return para.text.strip()

        # Fallback: first non-empty paragraph
        for para in doc.paragraphs:
            if para.text.strip():
                return para.text.strip()[:80]
    except Exception:
        pass
    return fallback


def _detect_format(filename: str, content: bytes) -> str:
    """Auto-detect document format from filename and magic bytes.

    Returns: 'pdf' or 'docx'
    """
    # 1. Extension
    if filename:
        ext = filename.rsplit(".", 1)[-1].lower() if "." in filename else ""
        if ext == "pdf":
            return "pdf"
        if ext in ("docx", "doc"):
            return "docx"

    # 2. Magic bytes
    if content[:5] == b"%PDF-":
        return "pdf"
    if content[:4] == b"PK\x03\x04":  # ZIP (DOCX is a ZIP archive)
        return "docx"

    # Default
    return "docx"


def _match_field_to_widget(field: dict, widgets: list[dict]) -> str | None:
    """Match a detected field to a PDF widget by value, name, or label.

    Returns widget_name if matched, None otherwise.
    """
    cv = (field.get("current_value") or "").strip()
    hint = (field.get("context_hint") or "").lower()
    field_name = field.get("field_name", "")

    # 1. Exact field_name match (LLM used widget_name directly as field_name)
    for w in widgets:
        if w["widget_name"] == field_name:
            logger.info("Widget match by name: %s → %s", field_name, w["widget_name"])
            return w["widget_name"]

    # 2. Exact value match (for filled fields like dates, NIF)
    if cv and not cv.startswith(".") and len(cv) > 1:
        for w in widgets:
            wv = (w.get("widget_value") or "").strip()
            if wv and wv == cv:
                logger.info("Widget match by value: %s → %s (cv='%s')",
                            field_name, w["widget_name"], cv[:20])
                return w["widget_name"]

    # 3. Case-insensitive name match
    fn_lower = field_name.lower()
    for w in widgets:
        wn_lower = w["widget_name"].lower()
        if wn_lower == fn_lower or wn_lower in fn_lower or fn_lower in wn_lower:
            logger.info("Widget match by name (fuzzy): %s → %s", field_name, w["widget_name"])
            return w["widget_name"]

    # 4. Label overlap (check if widget label words appear in context_hint)
    if hint:
        for w in widgets:
            wlabel = (w.get("label") or "").lower()
            if wlabel and len(wlabel) > 3:
                words = [word for word in wlabel.split() if len(word) > 3]
                if words and all(word in hint for word in words[:3]):
                    logger.info("Widget match by label: %s → %s", field_name, w["widget_name"])
                    return w["widget_name"]

    return None


@router.post("/analyze", response_model=AnalyzeResponse)
async def analyze_document(
    tenant_id: str = Form(...),
    user_id: str = Form(...),
    document_id: str = Form(None),
    user_intent: str = Form("modification"),
    max_fields: int = Form(30),
    file: UploadFile | None = File(None),
    _api_key: str = Depends(verify_api_key),
):
    """Detect variable fields in a document.

    Accepts either:
    - Direct file upload (multipart)
    - document_id to fetch from storage
    """
    settings = get_settings()
    file_bytes: bytes | None = None
    source_title = "document"

    # Path 1: Direct upload
    if file is not None:
        content = await file.read()
        if len(content) > settings.max_document_size_mb * 1024 * 1024:
            raise HTTPException(
                status_code=413,
                detail=f"File exceeds {settings.max_document_size_mb}MB limit",
            )
        file_bytes = content
        source_title = file.filename or "uploaded_document"

    # Path 2: Fetch by document_id
    elif document_id:
        weaviate = get_weaviate_client()
        metadata = await weaviate.get_document_metadata(document_id, tenant_id)
        if not metadata:
            raise HTTPException(status_code=404, detail="Document not found")

        file_path = metadata.get("file_path", "")
        if not file_path:
            raise HTTPException(status_code=404, detail="Document file path not found")

        source_title = metadata.get("file_name", metadata.get("title", "document"))

        storage = get_storage_client()
        file_bytes = await storage.download_document(file_path, tenant_id)
        if not file_bytes:
            raise HTTPException(status_code=404, detail="Could not download document")
    else:
        raise HTTPException(
            status_code=400, detail="Provide either a file upload or document_id"
        )

    # Auto-detect format
    source_format = _detect_format(source_title, file_bytes)
    pdf_has_widgets = False
    widget_map: dict[str, str] = {}  # field_name → widget_name

    # Extract text based on format
    if source_format == "pdf":
        try:
            document_text, title = extract_text_from_pdf(
                file_bytes, max_chars=settings.max_source_chars,
            )
        except Exception as e:
            raise HTTPException(status_code=422, detail=f"Could not parse PDF: {e}")

        # Check for AcroForm fields
        pdf_has_widgets = has_form_fields(file_bytes)
        if pdf_has_widgets:
            widgets = extract_widgets_from_pdf(file_bytes)
            logger.info("PDF has %d form widgets", len(widgets))
            # Append widget info to document text for LLM context
            widget_section = "\n\n--- EDITABLE FORM FIELDS ---\n"
            widget_section += "These are the fillable fields in the PDF form. Use widget_name as field_name.\n"
            widget_section += "For CheckBox fields, value should be 'true' or 'false'.\n"
            widget_section += "For Cifra fields (max_len=1), value is a single digit.\n\n"
            for w in widgets:
                extra = ""
                if w.get("max_len"):
                    extra = f" | max_len: {w['max_len']}"
                if w["widget_type"] == "CheckBox":
                    extra = " | values: true/false"
                widget_section += (
                    f"Widget: {w['widget_name']} | "
                    f"Label: {w['label']} | "
                    f"Value: \"{w['widget_value']}\" | "
                    f"Type: {w['widget_type']}{extra}\n"
                )
            document_text += widget_section
    else:
        try:
            document_text = _extract_text_from_docx(file_bytes)
        except Exception as e:
            raise HTTPException(status_code=422, detail=f"Could not parse DOCX: {e}")
        title = _extract_title_from_docx(file_bytes, source_title)

    if not document_text.strip():
        raise HTTPException(status_code=422, detail="Document contains no extractable text")

    # LLM analysis
    analyzer = get_analyzer()
    result = await analyzer.analyze(
        document_text=document_text,
        document_title=title,
        user_intent=user_intent,
        max_fields=max_fields,
    )

    # For PDF with widgets: map detected fields to widget names by matching
    # current_value or by label similarity
    if pdf_has_widgets:
        widgets = extract_widgets_from_pdf(file_bytes)
        for field in result.get("fields", []):
            matched_widget = _match_field_to_widget(field, widgets)
            if matched_widget:
                widget_map[field["field_name"]] = matched_widget
                field["widget_name"] = matched_widget

    # Create session and store source bytes
    store = get_session_store()
    session = await store.create_session(tenant_id, user_id)
    session.source_title = source_title
    session.source_format = source_format
    session.document_type = result.get("document_type", "")
    session.fields = result.get("fields", [])
    session.confidence = result.get("confidence", 0.0)
    session.status = SessionStatus.ANALYZED
    await store.update_session(session)

    # Store original file bytes for later rendering
    await store.store_blob(session.session_id, "source", file_bytes)

    return AnalyzeResponse(
        session_id=session.session_id,
        source_title=source_title,
        document_type=result.get("document_type", ""),
        fields=result.get("fields", []),
        confidence=result.get("confidence", 0.0),
    )
