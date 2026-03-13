"""DOCX template preparation — inserts Jinja2 markers at detected field positions.

This module handles the "split-run problem": Word's OOXML format arbitrarily
splits text across multiple <w:r> (run) elements within a paragraph. A simple
search like `if "01/01/2024" in run.text` fails because the date might be split
as ["01/0", "1/20", "24"] across three runs.

Strategy:
1. Concatenate all run texts in a paragraph to find the match position
2. Identify which runs span the match
3. Replace the matched text with {{field_name}}, preserving the first run's formatting
4. Clear text from subsequent partial runs
"""

import logging
import re
from io import BytesIO
from typing import Any

from docx import Document
from docx.table import Table
from docx.text.paragraph import Paragraph

logger = logging.getLogger(__name__)

# Only allow safe identifier characters in field names (prevents Jinja2 injection)
_SAFE_FIELD_NAME = re.compile(r"^[a-zA-Z_][a-zA-Z0-9_]*$")


class TemplatePreparer:
    """Inserts docxtpl-compatible {{markers}} into a DOCX at field positions."""

    def prepare(
        self,
        docx_bytes: bytes,
        fields: list[dict[str, Any]],
        fields_to_mark: list[str] | None = None,
        custom_fields: list[dict[str, Any]] | None = None,
    ) -> tuple[bytes, int, list[str]]:
        """Insert Jinja2 markers into a DOCX document.

        Args:
            docx_bytes: Original DOCX file bytes.
            fields: List of detected field dicts (from analyzer).
            fields_to_mark: Subset of field_names to mark. None = all fields.
            custom_fields: Additional custom fields with search_text.

        Returns:
            Tuple of (prepared_bytes, markers_inserted, failed_fields).
        """
        doc = Document(BytesIO(docx_bytes))

        # Build replacement map: search_text → marker_text
        replacements: list[tuple[str, str]] = []

        for field in fields:
            name = field.get("field_name", "")
            current_value = field.get("current_value", "")
            if not name or not current_value:
                continue
            if fields_to_mark is not None and name not in fields_to_mark:
                continue
            if not _SAFE_FIELD_NAME.match(name):
                logger.warning("Rejected unsafe field name: '%s'", name)
                continue
            replacements.append((current_value, "{{" + name + "}}"))

        # Custom fields: user-specified search text
        if custom_fields:
            for cf in custom_fields:
                search_text = cf.get("search_text", "")
                field_name = cf.get("field_name", "")
                if search_text and field_name:
                    if not _SAFE_FIELD_NAME.match(field_name):
                        logger.warning("Rejected unsafe custom field name: '%s'", field_name)
                        continue
                    replacements.append((search_text, "{{" + field_name + "}}"))

        # Sort by length descending to avoid partial matches
        replacements.sort(key=lambda x: len(x[0]), reverse=True)

        inserted = 0
        failed = []

        for search_text, marker in replacements:
            field_name = marker.strip("{}")
            count = self._replace_in_document(doc, search_text, marker)
            if count > 0:
                inserted += count
                logger.info("Inserted %d marker(s) for '%s'", count, field_name)
            else:
                failed.append(field_name)
                logger.warning("Could not find text for field '%s': '%s'", field_name, search_text[:50])

        # Save prepared document
        buf = BytesIO()
        doc.save(buf)
        return buf.getvalue(), inserted, failed

    def _replace_in_document(self, doc: Document, search_text: str, marker: str) -> int:
        """Replace all occurrences of search_text with marker across the entire document."""
        count = 0

        # Body paragraphs
        for para in doc.paragraphs:
            count += self._replace_in_paragraph(para, search_text, marker)

        # Tables (nested paragraphs in cells)
        for table in doc.tables:
            count += self._replace_in_table(table, search_text, marker)

        # Headers and footers
        for section in doc.sections:
            for header in [section.header, section.first_page_header, section.even_page_header]:
                if header and header.is_linked_to_previous is False:
                    for para in header.paragraphs:
                        count += self._replace_in_paragraph(para, search_text, marker)
                    for table in header.tables:
                        count += self._replace_in_table(table, search_text, marker)
            for footer in [section.footer, section.first_page_footer, section.even_page_footer]:
                if footer and footer.is_linked_to_previous is False:
                    for para in footer.paragraphs:
                        count += self._replace_in_paragraph(para, search_text, marker)
                    for table in footer.tables:
                        count += self._replace_in_table(table, search_text, marker)

        return count

    def _replace_in_table(self, table: Table, search_text: str, marker: str) -> int:
        """Replace in all cells of a table (including nested tables)."""
        count = 0
        for row in table.rows:
            for cell in row.cells:
                for para in cell.paragraphs:
                    count += self._replace_in_paragraph(para, search_text, marker)
                # Nested tables
                for nested_table in cell.tables:
                    count += self._replace_in_table(nested_table, search_text, marker)
        return count

    def _replace_in_paragraph(
        self, paragraph: Paragraph, search_text: str, marker: str
    ) -> int:
        """Replace search_text with marker in a paragraph, handling split runs.

        This is the core algorithm that solves the split-run problem:
        1. Concatenate all run texts to form the full paragraph text
        2. Find the search_text position in the concatenated string
        3. Map character positions back to individual runs
        4. Replace text across the spanning runs, preserving first run's formatting
        5. Rebuild run map and repeat until no more matches (iterative, not recursive)
        """
        count = 0
        max_iterations = 200  # Safety limit for pathological documents

        for _ in range(max_iterations):
            runs = paragraph.runs
            if not runs:
                break

            # Build concatenated text and position map
            full_text = ""
            run_map: list[tuple[int, int, int]] = []
            for i, run in enumerate(runs):
                start = len(full_text)
                run_text = run.text or ""
                full_text += run_text
                run_map.append((start, start + len(run_text), i))

            pos = full_text.find(search_text)
            if pos == -1:
                break

            match_end = pos + len(search_text)

            # Find which runs span this match
            spanning_runs = []
            for start, end, run_idx in run_map:
                if start < match_end and end > pos:
                    spanning_runs.append((start, end, run_idx))

            if not spanning_runs:
                break

            # Replace across spanning runs
            self._replace_across_runs(
                runs, spanning_runs, pos, match_end, marker
            )
            count += 1

        return count

    def _replace_across_runs(
        self,
        runs: list,
        spanning_runs: list[tuple[int, int, int]],
        match_start: int,
        match_end: int,
        marker: str,
    ) -> None:
        """Replace matched text across multiple runs with the marker.

        The first spanning run gets the marker text (preserving its formatting).
        Subsequent spanning runs have the matched portion of their text removed.
        """
        for i, (run_start, run_end, run_idx) in enumerate(spanning_runs):
            run = runs[run_idx]
            run_text = run.text or ""

            # Calculate overlap between this run and the match
            overlap_start = max(match_start, run_start) - run_start
            overlap_end = min(match_end, run_end) - run_start

            if i == 0:
                # First run: replace matched portion with marker
                before = run_text[:overlap_start]
                after = run_text[overlap_end:]
                run.text = before + marker + after
            else:
                # Subsequent runs: remove matched portion
                before = run_text[:overlap_start]
                after = run_text[overlap_end:]
                run.text = before + after


_preparer = None


def get_template_preparer() -> TemplatePreparer:
    global _preparer
    if _preparer is None:
        _preparer = TemplatePreparer()
    return _preparer
