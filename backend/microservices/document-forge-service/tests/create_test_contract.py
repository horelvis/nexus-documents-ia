"""
Generate a test DOCX contract with all 8 ForgeFieldType values.

Creates a realistic employment contract (contrato laboral) containing
at least one field of each type: text, date, number, currency,
name, address, email, phone.

Usage:
    python create_test_contract.py [output_path]
    Default output: ~/contrato-test-forge.docx
"""

import sys
import os
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT


def create_test_contract(output_path: str):
    doc = Document()

    # -- Styles --
    style = doc.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(11)
    style.paragraph_format.space_after = Pt(6)

    # -- Title --
    title = doc.add_heading("CONTRATO DE TRABAJO TEMPORAL", level=1)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER

    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = subtitle.add_run("Ref: CT-2026-00142")
    run.font.size = Pt(10)
    run.font.color.rgb = RGBColor(128, 128, 128)

    doc.add_paragraph()  # spacer

    # -- Section: REUNIDOS --
    doc.add_heading("REUNIDOS", level=2)

    # NAME fields (employer + employee)
    doc.add_paragraph(
        "De una parte, D./Dña. María García López, en calidad de Directora de Recursos Humanos, "
        "en nombre y representación de la empresa Soluciones Tecnológicas del Mediterráneo S.L., "
        "con CIF B-12345678 y domicilio social en la dirección que se indica más abajo."
    )

    doc.add_paragraph(
        "De otra parte, D./Dña. Carlos Martínez Rodríguez, mayor de edad, con DNI 48765432-Z, "
        "número de afiliación a la Seguridad Social 281234567890, y domicilio en la dirección "
        "que se indica más abajo."
    )

    doc.add_paragraph(
        "Ambas partes se reconocen mutuamente la capacidad legal necesaria para el otorgamiento "
        "del presente contrato de trabajo temporal, y a tal efecto:"
    )

    # -- Section: MANIFIESTAN --
    doc.add_heading("MANIFIESTAN", level=2)

    doc.add_paragraph(
        "Que la empresa necesita cubrir temporalmente el puesto de Desarrollador Full Stack Senior "
        "en el departamento de Ingeniería de Software, debido al incremento de la actividad "
        "productiva derivado del proyecto europeo H2020-DIGITECH."
    )

    # -- Section: CLÁUSULAS --
    doc.add_heading("CLÁUSULAS", level=2)

    # Clause 1: TEXT field (job title / department)
    doc.add_heading("PRIMERA — Objeto del contrato", level=3)
    doc.add_paragraph(
        "El/La trabajador/a es contratado/a para desempeñar las funciones propias del puesto "
        "de Desarrollador Full Stack Senior, categoría profesional Grupo 2 — Ingenieros Técnicos, "
        "dentro del departamento de Ingeniería de Software."
    )

    # Clause 2: DATE fields (start date, end date)
    doc.add_heading("SEGUNDA — Duración y período de prueba", level=3)
    doc.add_paragraph(
        "El presente contrato tendrá una duración determinada, comenzando el día 15/03/2026 "
        "y finalizando el día 14/09/2026, salvo prórroga expresa acordada por ambas partes."
    )
    doc.add_paragraph(
        "Se establece un período de prueba de 2 meses, durante el cual cualquiera de las "
        "partes podrá resolver la relación laboral sin necesidad de preaviso."
    )

    # Clause 3: NUMBER field (hours/week) + CURRENCY field (salary)
    doc.add_heading("TERCERA — Jornada laboral y retribución", level=3)
    doc.add_paragraph(
        "La jornada laboral será de 40 horas semanales, distribuidas de lunes a viernes "
        "en horario de 09:00 a 18:00 con una hora de descanso para la comida."
    )
    doc.add_paragraph(
        "El/La trabajador/a percibirá una retribución bruta anual de 42.000,00 EUR, "
        "distribuida en 14 pagas (12 mensualidades + 2 pagas extraordinarias en junio y diciembre). "
        "El salario mensual bruto será de 3.000,00 EUR."
    )

    # Clause 4: ADDRESS field (workplace)
    doc.add_heading("CUARTA — Lugar de trabajo", level=3)
    doc.add_paragraph(
        "El/La trabajador/a prestará sus servicios en el centro de trabajo ubicado en "
        "Calle de la Innovación 42, 3ª planta, 30100 Espinardo, Murcia, España, "
        "sin perjuicio de los desplazamientos que pudieran ser necesarios para el "
        "desarrollo de su actividad profesional."
    )

    # Clause 5: Additional terms (TEXT)
    doc.add_heading("QUINTA — Vacaciones y permisos", level=3)
    doc.add_paragraph(
        "El/La trabajador/a tendrá derecho a 23 días laborables de vacaciones anuales, "
        "que se disfrutarán de acuerdo con lo establecido en el Convenio Colectivo "
        "del sector de Consultoría y Empresas de Tecnología."
    )

    # Clause 6: Confidentiality (TEXT)
    doc.add_heading("SEXTA — Confidencialidad", level=3)
    doc.add_paragraph(
        "El/La trabajador/a se compromete a mantener la más estricta confidencialidad "
        "sobre toda la información a la que tenga acceso durante la vigencia del contrato, "
        "así como a no utilizar dicha información para fines ajenos a la empresa."
    )

    # -- Section: Contact Table with EMAIL + PHONE fields --
    doc.add_heading("DATOS DE CONTACTO", level=2)
    doc.add_paragraph("A efectos de comunicaciones, las partes facilitan los siguientes datos:")

    table = doc.add_table(rows=3, cols=3, style="Table Grid")
    table.alignment = WD_TABLE_ALIGNMENT.CENTER

    # Header row
    headers = ["", "Empresa", "Trabajador/a"]
    for i, h in enumerate(headers):
        cell = table.rows[0].cells[i]
        cell.text = h
        for p in cell.paragraphs:
            for run in p.runs:
                run.font.bold = True
                run.font.size = Pt(10)

    # EMAIL row
    table.rows[1].cells[0].text = "Email"
    table.rows[1].cells[1].text = "rrhh@soluciones-med.es"
    table.rows[1].cells[2].text = "carlos.martinez@gmail.com"

    # PHONE row
    table.rows[2].cells[0].text = "Teléfono"
    table.rows[2].cells[1].text = "+34 968 123 456"
    table.rows[2].cells[2].text = "+34 612 345 678"

    doc.add_paragraph()  # spacer

    # -- Section: FIRMAS --
    doc.add_heading("FIRMAS", level=2)

    # Signing date (another DATE)
    doc.add_paragraph(
        "Y para que así conste, ambas partes firman el presente contrato por duplicado "
        "y a un solo efecto, en Murcia, a 10 de marzo de 2026."
    )

    doc.add_paragraph()

    # Signature table
    sig_table = doc.add_table(rows=2, cols=2, style="Table Grid")
    sig_table.alignment = WD_TABLE_ALIGNMENT.CENTER

    sig_table.rows[0].cells[0].text = "Por la empresa:"
    sig_table.rows[0].cells[1].text = "El/La trabajador/a:"

    sig_table.rows[1].cells[0].text = "\n\n\nFdo: María García López\nDirectora de RRHH"
    sig_table.rows[1].cells[1].text = "\n\n\nFdo: Carlos Martínez Rodríguez\nDNI: 48765432-Z"

    # -- Footer info --
    doc.add_paragraph()
    footer_para = doc.add_paragraph()
    footer_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = footer_para.add_run(
        "Soluciones Tecnológicas del Mediterráneo S.L. — CIF: B-12345678\n"
        "Calle de la Innovación 42, 30100 Espinardo, Murcia\n"
        "Tel: +34 968 123 456 — Email: info@soluciones-med.es"
    )
    run.font.size = Pt(8)
    run.font.color.rgb = RGBColor(150, 150, 150)

    # Save
    doc.save(output_path)
    print(f"Created: {output_path}")
    print(f"Size: {os.path.getsize(output_path):,} bytes")
    print()
    print("Field types included:")
    print("  NAME     → María García López, Carlos Martínez Rodríguez")
    print("  DATE     → 15/03/2026, 14/09/2026, 10 de marzo de 2026")
    print("  TEXT     → Desarrollador Full Stack Senior, Ingeniería de Software")
    print("  NUMBER   → 40 (horas semanales), 23 (días vacaciones), 2 (meses prueba)")
    print("  CURRENCY → 42.000,00 EUR, 3.000,00 EUR")
    print("  ADDRESS  → Calle de la Innovación 42, 3ª planta, 30100 Espinardo, Murcia")
    print("  EMAIL    → rrhh@soluciones-med.es, carlos.martinez@gmail.com")
    print("  PHONE    → +34 968 123 456, +34 612 345 678")


if __name__ == "__main__":
    output = sys.argv[1] if len(sys.argv) > 1 else os.path.expanduser("~/contrato-test-forge.docx")
    create_test_contract(output)
