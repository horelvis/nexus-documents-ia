"""
DOCX renderer for verified generation reports.

Uses python-docx to programmatically generate Word documents
from verified generation sessions. This avoids template complexity
and gives full control over styling.
"""

import logging
from io import BytesIO
from typing import Any, Dict

logger = logging.getLogger(__name__)


class DocxRenderer:
    """Render verified generation results as DOCX."""

    def render_verified_report(self, data: Dict[str, Any]) -> bytes:
        """
        Render a verified document report as DOCX bytes.

        Args:
            data: Dict with keys: query, session_id, created_at, document_text,
                  claims, claims_verified, claims_corrected, claims_rejected,
                  average_confidence, execution_time_ms, sources (optional)

        Returns:
            DOCX file content as bytes.
        """
        from docx import Document
        from docx.shared import Pt, Cm, RGBColor
        from docx.enum.text import WD_ALIGN_PARAGRAPH
        from docx.enum.table import WD_TABLE_ALIGNMENT

        doc = Document()

        # Page margins
        for section in doc.sections:
            section.top_margin = Cm(2)
            section.bottom_margin = Cm(2)
            section.left_margin = Cm(2.5)
            section.right_margin = Cm(2.5)

        confidence_pct = round((data.get("average_confidence", 0) or 0) * 100)

        # ── Title ──
        title = doc.add_heading("Informe de Verificación", level=1)
        for run in title.runs:
            run.font.color.rgb = RGBColor(0x1E, 0x29, 0x3B)

        # ── Metadata ──
        session_id = data.get("session_id", "") or ""
        for label, value in [
            ("Tema: ", data.get("query", "N/A")),
            ("Fecha: ", data.get("created_at", "N/A")),
            ("Sesión: ", f"{session_id[:16]}…"),
        ]:
            p = doc.add_paragraph()
            p.paragraph_format.space_after = Pt(2)
            r = p.add_run(label)
            r.bold = True
            r.font.size = Pt(9)
            r.font.color.rgb = RGBColor(0x47, 0x55, 0x69)
            r2 = p.add_run(value)
            r2.font.size = Pt(9)

        # ── Source document info ──
        source_filenames = data.get("source_filenames", [])
        if source_filenames:
            p = doc.add_paragraph()
            p.paragraph_format.space_after = Pt(2)
            label = "Documento fuente: " if len(source_filenames) == 1 else "Documentos fuente: "
            r = p.add_run(label)
            r.bold = True
            r.font.size = Pt(9)
            r.font.color.rgb = RGBColor(0x47, 0x55, 0x69)
            r2 = p.add_run(", ".join(source_filenames))
            r2.font.size = Pt(9)

        source_summary = data.get("source_summary", "")
        if source_summary:
            p = doc.add_paragraph()
            p.paragraph_format.space_after = Pt(6)
            r = p.add_run("Resumen del documento fuente")
            r.bold = True
            r.font.size = Pt(8)
            r.font.color.rgb = RGBColor(0x25, 0x63, 0xEB)
            p2 = doc.add_paragraph()
            p2.paragraph_format.space_after = Pt(8)
            r2 = p2.add_run(source_summary)
            r2.font.size = Pt(9)
            r2.font.color.rgb = RGBColor(0x47, 0x55, 0x69)

        # ── Stats table ──
        stats = doc.add_table(rows=2, cols=4)
        stats.alignment = WD_TABLE_ALIGNMENT.CENTER
        stats.style = "Light Grid Accent 1"

        stat_headers = ["Verificados", "Corregidos", "Rechazados", "Confianza"]
        stat_values = [
            str(data.get("claims_verified", 0)),
            str(data.get("claims_corrected", 0)),
            str(data.get("claims_rejected", 0)),
            f"{confidence_pct}%",
        ]
        for i, (h, v) in enumerate(zip(stat_headers, stat_values)):
            for row_idx, text, size in [(0, h, Pt(8)), (1, v, Pt(14))]:
                cell = stats.rows[row_idx].cells[i]
                cell.text = ""
                p = cell.paragraphs[0]
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                run = p.add_run(text)
                run.bold = True
                run.font.size = size

        doc.add_paragraph()

        # ── Document Body ──
        doc.add_heading("Documento Verificado", level=2)
        document_text = data.get("document_text", "")
        for para_text in document_text.split("\n"):
            if para_text.strip():
                p = doc.add_paragraph(para_text.strip())
                p.paragraph_format.space_after = Pt(6)

        doc.add_paragraph()

        # ── Claims Detail ──
        claims = data.get("claims", [])
        doc.add_heading(f"Detalle de Claims ({len(claims)})", level=2)

        if claims:
            ct = doc.add_table(rows=1 + len(claims), cols=4)
            ct.style = "Table Grid"

            # Header
            for i, h in enumerate(["#", "Claim", "Estado", "Fuentes"]):
                ct.rows[0].cells[i].text = ""
                run = ct.rows[0].cells[i].paragraphs[0].add_run(h)
                run.bold = True
                run.font.size = Pt(9)

            # Data rows
            status_labels = {
                "verified": "VERIFICADO",
                "corrected": "CORREGIDO",
                "rejected": "RECHAZADO",
            }
            status_colors = {
                "verified": RGBColor(0x16, 0x6F, 0x34),
                "corrected": RGBColor(0x92, 0x40, 0x0E),
                "rejected": RGBColor(0x99, 0x1B, 0x1B),
            }
            source_type_labels = {
                "web": "Web",
                "public_knowledge": "Web",
                "uploaded": "Cargado",
                "internal": "Interno",
                "jurisprudence": "Jurisp.",
                "doi": "DOI",
                "crossref": "CrossRef",
                "doi_invalid": "DOI Inv.",
                "doi_mismatch": "DOI ≠",
                "citation_unverified": "Cita ?",
            }
            for idx, claim in enumerate(claims):
                row = ct.rows[idx + 1]
                status = claim.get("status", "verified")
                conf = round((claim.get("confidence", 0) or 0) * 100)

                # Number
                row.cells[0].text = ""
                run = row.cells[0].paragraphs[0].add_run(str(idx + 1))
                run.font.size = Pt(9)
                run.bold = True

                # Text
                row.cells[1].text = ""
                run = row.cells[1].paragraphs[0].add_run(claim.get("text", ""))
                run.font.size = Pt(9)

                # Status + verification type
                row.cells[2].text = ""
                p = row.cells[2].paragraphs[0]
                label = status_labels.get(status, status.upper())
                run = p.add_run(f"{label} ({conf}%)")
                run.font.size = Pt(9)
                run.bold = True
                run.font.color.rgb = status_colors.get(status, RGBColor(0x47, 0x55, 0x69))
                vtype = claim.get("verification_type")
                if vtype:
                    vtype_labels = {
                        "corroborated": "Corroborado",
                        "fidelity_only": "Fidelidad",
                        "independent": "Verificado",
                    }
                    vtype_label = vtype_labels.get(vtype, vtype)
                    run2 = p.add_run(f"\n{vtype_label}")
                    run2.font.size = Pt(8)
                    run2.italic = True
                    run2.font.color.rgb = RGBColor(0x64, 0x74, 0x8B)

                # Verification reason
                vreason = claim.get("verification_reason")
                if vreason:
                    reason_run = p.add_run(f"\n{vreason}")
                    reason_run.font.size = Pt(8)
                    reason_run.italic = True
                    reason_run.font.color.rgb = RGBColor(0x47, 0x55, 0x69)

                # Sources
                row.cells[3].text = ""
                evidence_sources = claim.get("evidence_sources", [])
                if evidence_sources:
                    p = row.cells[3].paragraphs[0]
                    for si, src in enumerate(evidence_sources):
                        src_type = source_type_labels.get(src.get("source", "internal"), "Interno")
                        title = src.get("title") or src.get("id", "")[:20]
                        line = f"[{src_type}] {title}"
                        run = p.add_run(line)
                        run.font.size = Pt(8)
                        run.font.color.rgb = RGBColor(0x64, 0x74, 0x8B)
                        if si < len(evidence_sources) - 1:
                            p.add_run("\n").font.size = Pt(8)

        doc.add_paragraph()

        # ── Sources ──
        sources_data = data.get("sources", [])
        if sources_data:
            doc.add_heading(f"Fuentes Consultadas ({len(sources_data)})", level=2)

            st = doc.add_table(rows=1 + len(sources_data), cols=3)
            st.style = "Table Grid"

            for i, h in enumerate(["Tipo", "Título", "URL"]):
                st.rows[0].cells[i].text = ""
                run = st.rows[0].cells[i].paragraphs[0].add_run(h)
                run.bold = True
                run.font.size = Pt(9)

            type_labels = {
                "web": "Web", "public_knowledge": "Web",
                "uploaded": "Cargado", "internal": "Interno",
                "doi": "DOI Validado", "crossref": "CrossRef",
                "doi_invalid": "DOI Inválido", "doi_mismatch": "DOI ≠",
                "citation_unverified": "Cita No Verif.",
            }
            for idx, src in enumerate(sources_data):
                row = st.rows[idx + 1]
                src_type = src.get("source", "internal")

                row.cells[0].text = ""
                row.cells[0].paragraphs[0].add_run(
                    type_labels.get(src_type, "Interno")
                ).font.size = Pt(9)

                row.cells[1].text = ""
                row.cells[1].paragraphs[0].add_run(
                    src.get("title") or src.get("id", "")
                ).font.size = Pt(9)

                row.cells[2].text = ""
                row.cells[2].paragraphs[0].add_run(
                    src.get("url", "") or "—"
                ).font.size = Pt(9)

        # ── Validación Bibliográfica (DOI) ──
        doi_validations = data.get("doi_validations", [])
        if doi_validations:
            valid_count = sum(1 for d in doi_validations if d.get("valid"))
            doc.add_heading(
                f"Validación Bibliográfica ({valid_count}/{len(doi_validations)} DOIs válidos)",
                level=2,
            )
            dt = doc.add_table(rows=1 + len(doi_validations), cols=3)
            dt.style = "Table Grid"

            for i, h in enumerate(["Estado", "DOI", "Publicación"]):
                dt.rows[0].cells[i].text = ""
                run = dt.rows[0].cells[i].paragraphs[0].add_run(h)
                run.bold = True
                run.font.size = Pt(9)

            for idx, dv in enumerate(doi_validations):
                row = dt.rows[idx + 1]
                is_valid = dv.get("valid", False)
                meta = dv.get("metadata", {}) or {}

                # Estado
                row.cells[0].text = ""
                status_run = row.cells[0].paragraphs[0].add_run(
                    "Válido" if is_valid else "Inválido"
                )
                status_run.font.size = Pt(9)
                status_run.bold = True
                status_run.font.color.rgb = (
                    RGBColor(0x16, 0x6F, 0x34) if is_valid
                    else RGBColor(0x99, 0x1B, 0x1B)
                )

                # DOI
                row.cells[1].text = ""
                doi_run = row.cells[1].paragraphs[0].add_run(dv.get("doi", ""))
                doi_run.font.size = Pt(8)
                if is_valid:
                    doi_run.font.color.rgb = RGBColor(0x25, 0x63, 0xEB)
                else:
                    doi_run.font.color.rgb = RGBColor(0x99, 0x1B, 0x1B)

                # Publicación
                row.cells[2].text = ""
                if meta.get("title"):
                    title_text = meta["title"]
                    authors = meta.get("authors", [])
                    year = meta.get("year", "")
                    if authors:
                        title_text += f" — {', '.join(authors[:3])}"
                    if year:
                        title_text += f" ({year})"
                    pub_run = row.cells[2].paragraphs[0].add_run(title_text)
                    pub_run.font.size = Pt(8)
                else:
                    pub_run = row.cells[2].paragraphs[0].add_run("No se pudo resolver")
                    pub_run.font.size = Pt(8)
                    pub_run.italic = True
                    pub_run.font.color.rgb = RGBColor(0x94, 0xA3, 0xB8)

            doc.add_paragraph()

        # ── Metodología ──
        doc.add_heading("Metodología de Verificación", level=2)

        intro = doc.add_paragraph()
        intro.paragraph_format.space_after = Pt(6)
        r = intro.add_run(
            "Cada afirmación se evalúa en dos niveles independientes, siguiendo principios de "
            "SAFE (Google DeepMind, 2024), CoVe (Meta, 2024) y FFCI (Koto et al.):"
        )
        r.font.size = Pt(9)
        r.font.color.rgb = RGBColor(0x47, 0x55, 0x69)

        mt = doc.add_table(rows=4, cols=3)
        mt.style = "Table Grid"

        for i, h in enumerate(["Tipo", "Significado", "Conf. máx."]):
            mt.rows[0].cells[i].text = ""
            run = mt.rows[0].cells[i].paragraphs[0].add_run(h)
            run.bold = True
            run.font.size = Pt(9)

        method_rows = [
            (
                "Fidelidad",
                "La afirmación representa fielmente el documento fuente "
                "(NLI entailment), pero no se encontró evidencia independiente.",
                "80%",
                RGBColor(0x92, 0x40, 0x0E),
            ),
            (
                "Corroborado",
                "Fiel al documento fuente y además respaldada por evidencia "
                "independiente (otros documentos, legislación, web).",
                "100%",
                RGBColor(0x16, 0x6F, 0x34),
            ),
            (
                "Verificado",
                "Sin documento fuente. Respaldada exclusivamente por "
                "evidencia independiente externa.",
                "100%",
                RGBColor(0x1E, 0x40, 0xAF),
            ),
        ]
        for ri, (tipo, significado, conf, color) in enumerate(method_rows):
            row = mt.rows[ri + 1]
            row.cells[0].text = ""
            r = row.cells[0].paragraphs[0].add_run(tipo)
            r.font.size = Pt(9)
            r.bold = True
            r.font.color.rgb = color

            row.cells[1].text = ""
            r = row.cells[1].paragraphs[0].add_run(significado)
            r.font.size = Pt(8)

            row.cells[2].text = ""
            p = row.cells[2].paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            r = p.add_run(conf)
            r.font.size = Pt(9)
            r.bold = True
            r.font.color.rgb = color

        doc.add_paragraph()
        tiers = doc.add_paragraph()
        tiers.paragraph_format.space_after = Pt(4)
        r = tiers.add_run("Tier 1 — Fidelidad: ")
        r.bold = True
        r.font.size = Pt(8)
        r.font.color.rgb = RGBColor(0x64, 0x74, 0x8B)
        r = tiers.add_run(
            "Evaluación NLI que verifica si la afirmación está contenida (entailed) "
            "en el texto fuente, basada en FACTS Grounding (Google DeepMind, 2024). "
            "Confianza limitada al 80% porque la fuente de generación y verificación coinciden. "
        )
        r.font.size = Pt(8)
        r.font.color.rgb = RGBColor(0x64, 0x74, 0x8B)
        r = tiers.add_run("Tier 2 — Corroboración externa: ")
        r.bold = True
        r.font.size = Pt(8)
        r.font.color.rgb = RGBColor(0x64, 0x74, 0x8B)
        r = tiers.add_run(
            "Búsqueda en fuentes independientes (excluyendo documentos fuente), "
            "siguiendo SAFE (Google DeepMind, 2024) y CoVe (Meta, 2024)."
        )
        r.font.size = Pt(8)
        r.font.color.rgb = RGBColor(0x64, 0x74, 0x8B)

        # ── Footer ──
        doc.add_paragraph()
        footer = doc.add_paragraph()
        footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
        exec_ms = data.get("execution_time_ms", 0)
        run = footer.add_run(
            f"Generado por NouxCubeIA — Tiempo de ejecución: {exec_ms}ms"
        )
        run.font.size = Pt(8)
        run.font.color.rgb = RGBColor(0x94, 0xA3, 0xB8)

        # Write to bytes
        buf = BytesIO()
        doc.save(buf)
        docx_bytes = buf.getvalue()
        logger.info(
            f"DOCX rendered: {len(docx_bytes)} bytes for session "
            f"{session_id[:16]}"
        )
        return docx_bytes


_renderer = None


def get_docx_renderer() -> DocxRenderer:
    """Get singleton DocxRenderer instance."""
    global _renderer
    if _renderer is None:
        _renderer = DocxRenderer()
    return _renderer
