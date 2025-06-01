import logging
import os
import uuid
import datetime
import io
from typing import List, Dict, Any, Optional, BinaryIO, Union
from fastapi import UploadFile, HTTPException
import PyPDF2
from docx import Document as DocxDocument
import csv
import openpyxl
import chardet
from sqlalchemy.orm import Session, joinedload

from app.core.config import settings
from app.db.models import Document, DocumentChunk, Tag
from app.db.database import SessionLocal
from app.schemas.enums import IndexingStatus
from app.services.storage_service import StorageService
from app.services.embedding_service import EmbeddingService
from app.services.llm_service import LLMService

logger = logging.getLogger(__name__)


class DocumentService:
    """Servicio para gestión de documentos"""
    
    def __init__(self, tenant_id: str = None, user_id: str = None):
        """
        Inicializa el servicio de documentos.
        
        Args:
            tenant_id: ID del tenant (debe ser un UUID válido)
            user_id: ID del usuario actual
        """
        # Si tenant_id es None o el string "default", obtener el UUID real del tenant por defecto
        if not tenant_id or tenant_id == settings.DEFAULT_TENANT:
            from app.db.database import SessionLocal
            from app.db.models import Tenant
            db = SessionLocal()
            try:
                default_tenant = db.query(Tenant).filter(Tenant.name == settings.DEFAULT_TENANT).first()
                if default_tenant:
                    self.tenant_id = str(default_tenant.id)
                else:
                    raise ValueError(f"Default tenant '{settings.DEFAULT_TENANT}' not found in database")
            finally:
                db.close()
        else:
            self.tenant_id = tenant_id
            
        self.user_id = user_id
        self.storage_service = StorageService(self.tenant_id)
        self.embedding_service = EmbeddingService(self.tenant_id)
        self.llm_service = LLMService()

    async def _validate_file(self, file: UploadFile, filename: str) -> tuple[str, bytes, int]:
        """
        Validates the uploaded file, checks its extension and size.
        Returns the file extension, its contents as bytes, and its size.
        """
        if not file:
            raise HTTPException(status_code=400, detail="Archivo no proporcionado")

        contents = await file.read()
        await file.seek(0) 

        file_size = len(contents)
        file_ext = os.path.splitext(filename)[1][1:].lower() if "." in filename else ""

        if not file_ext or file_ext not in settings.ALLOWED_EXTENSIONS:
            raise HTTPException(
                status_code=400,
                detail=f"Tipo de archivo no permitido. Permitidos: {', '.join(settings.ALLOWED_EXTENSIONS)}"
            )
        
        if file_size == 0:
             raise HTTPException(status_code=400, detail="El archivo está vacío.")

        if file_size > settings.MAX_UPLOAD_SIZE:
            raise HTTPException(
                status_code=400,
                detail=f"Tamaño de archivo excede el límite de {settings.MAX_UPLOAD_SIZE // (1024*1024)}MB"
            )
        
        return file_ext, contents, file_size

    def _create_document_record(
        self, 
        db: Session, 
        title: str, 
        description: Optional[str], 
        filename: str, 
        file_ext: str, 
        file_size: int, 
        tags: Optional[List[str]]
    ) -> Document:
        """
        Creates the Document ORM object, generates doc_id, constructs file_path,
        and handles tags. Adds to session but does not commit.
        """
        doc_id = str(uuid.uuid4())
        file_path = f"documents/{doc_id}/{filename}"

        db_document = Document(
            id=doc_id,
            title=title,
            description=description or "",
            filename=filename,
            file_path=file_path,
            file_type=file_ext,
            file_size=file_size,
            tenant_id=self.tenant_id,
            created_by=self.user_id,
            indexed=IndexingStatus.PROCESSING
        )

        if tags:
            for tag_name in tags:
                if tag_name:
                    tag = db.query(Tag).filter(
                        Tag.name == tag_name,
                        Tag.tenant_id == self.tenant_id
                    ).first()
                    if not tag:
                        tag = Tag(name=tag_name, tenant_id=self.tenant_id)
                        db.add(tag)
                        db.flush()
                    db_document.tags.append(tag)
        
        db.add(db_document)
        db.flush()
        db.refresh(db_document)
        return db_document

    def _upload_file_to_storage(
        self, 
        file_contents: bytes, 
        file_path: str, 
        doc_id: str, 
        title: str, 
        file_ext: str
    ):
        """
        Uploads the file content to the storage service.
        """
        file_obj = io.BytesIO(file_contents)
        storage_metadata = {
            "doc_id": doc_id,
            "title": title,
            "content_type": f"application/{file_ext}"
        }
        
        upload_success = self.storage_service.upload_file(
            file=file_obj,
            object_name=file_path,
            metadata=storage_metadata
        )
        
        if not upload_success:
            raise Exception("Failed to store document file in cloud storage.")

    def _extract_and_index_text(
        self, 
        db: Session, 
        db_document: Document, 
        file_contents: bytes, 
        file_ext: str, 
        title: str
    ):
        """
        Extracts text, creates chunks, generates embeddings, and updates 
        the document's indexed status. Adds to session but does not commit.
        """
        file_obj_for_text = io.BytesIO(file_contents)
        document_text = self._extract_text(file_obj_for_text, file_ext)
        
        if document_text:
            chunks_data = self.embedding_service.chunk_text(document_text)
            for i, chunk_data in enumerate(chunks_data):
                chunk = DocumentChunk(
                    document_id=db_document.id,
                    chunk_index=i,
                    content=chunk_data["text"]
                )
                db.add(chunk)
            
            db.flush()

            indexing_success = self.embedding_service.add_document(
                doc_id=str(db_document.id),
                text=document_text,
                metadata={
                    "doc_id": str(db_document.id),
                    "title": title,
                    "file_type": file_ext,
                    "tenant_id": self.tenant_id
                }
            )
            db_document.indexed = IndexingStatus.INDEXED if indexing_success else IndexingStatus.INDEXING_ERROR
        else:
            db_document.indexed = IndexingStatus.INDEXING_ERROR
    
    async def process_document(
        self, 
        file: UploadFile, 
        title: str,
        description: Optional[str] = None,
        tags: Optional[List[str]] = None
    ) -> Dict[str, Any]:
        """
        Procesa un documento: lo valida, crea el registro en BD, lo almacena en GCS, 
        extrae texto, genera embeddings y lo indexa.
        """
        db = SessionLocal()
        try:
            filename = file.filename
            # 1. Validate file
            file_ext, file_contents, file_size = await self._validate_file(file, filename)

            # 2. Create document record in DB
            db_document = self._create_document_record(
                db=db,
                title=title,
                description=description,
                filename=filename,
                file_ext=file_ext,
                file_size=file_size,
                tags=tags
            )
            
            doc_id_str = str(db_document.id)
            file_path_str = db_document.file_path

            # 3. Upload file to storage
            self._upload_file_to_storage(
                file_contents=file_contents,
                file_path=file_path_str,
                doc_id=doc_id_str,
                title=db_document.title,
                file_ext=db_document.file_type
            )

            # 4. Extract text and index
            self._extract_and_index_text(
                db=db,
                db_document=db_document,
                file_contents=file_contents,
                file_ext=db_document.file_type,
                title=db_document.title
            )
            
            # 5. Commit all DB changes
            db.commit()
            db.refresh(db_document)
            
            # 6. Format and return response
            return {
                "id": str(db_document.id),
                "title": db_document.title,
                "description": db_document.description,
                "filename": db_document.filename,
                "file_type": db_document.file_type,
                "file_size": db_document.file_size,
                "indexed": db_document.indexed,
                "created_at": db_document.created_at.isoformat(),
                "tags": [tag.name for tag in db_document.tags]
            }
            
        except HTTPException:
            db.rollback()
            raise
        except Exception as e:
            db.rollback()
            logger.exception(f"Error processing document: {str(e)}")
            raise HTTPException(status_code=500, detail="An unexpected error occurred while processing the document.")
        finally:
            db.close()
    
    def get_document(self, doc_id: str) -> Dict[str, Any]:
        """
        Obtiene información detallada de un documento.
        """
        db = SessionLocal()
        
        try:
            document = db.query(Document).options(
                joinedload(Document.tags)
            ).filter(
                Document.id == doc_id,
                Document.tenant_id == self.tenant_id
            ).first()
            
            if not document:
                raise HTTPException(status_code=404, detail="Document not found")
            
            # Obtener primeros chunks para vista previa
            chunks = db.query(DocumentChunk).filter(
                DocumentChunk.document_id == doc_id
            ).order_by(
                DocumentChunk.chunk_index
            ).limit(3).all()
            
            chunks_dict = []
            for chunk in chunks:
                chunks_dict.append({
                    "id": chunk.id,
                    "document_id": str(chunk.document_id),
                    "chunk_index": chunk.chunk_index,
                    "content": chunk.content
                })
            
            result = {
                "id": str(document.id),
                "title": document.title,
                "description": document.description,
                "filename": document.filename,
                "file_path": document.file_path,
                "file_type": document.file_type,
                "file_size": document.file_size,
                "indexed": document.indexed,
                "created_at": document.created_at.isoformat(),
                "updated_at": document.updated_at.isoformat(),
                "tags": [tag.name for tag in document.tags],
                "preview_chunks": chunks_dict
            }
            
            return result
            
        except HTTPException:
            raise
        except Exception as e:
            logger.exception(f"Error getting document {doc_id}: {str(e)}")
            raise HTTPException(status_code=500, detail="An unexpected error occurred while retrieving the document.")
        finally:
            db.close()
    
    def delete_document(self, doc_id: str) -> Dict[str, Any]:
        """
        Elimina un documento y todos sus datos asociados.
        """
        db = SessionLocal()
        
        try:
            document = db.query(Document).filter(
                Document.id == doc_id,
                Document.tenant_id == self.tenant_id
            ).first()
            
            if not document:
                raise HTTPException(status_code=404, detail="Document not found")
            
            # Eliminar archivo del almacenamiento
            self.storage_service.delete_file(document.file_path)
            
            # Eliminar del vector store
            self.embedding_service.delete_document(doc_id)
            
            # Eliminar de la base de datos
            db.delete(document)
            db.commit()
            
            return {"message": f"Document {doc_id} deleted successfully"}
            
        except HTTPException:
            raise
        except Exception as e:
            db.rollback()
            logger.exception(f"Error deleting document {doc_id}: {str(e)}")
            raise HTTPException(status_code=500, detail="An unexpected error occurred while deleting the document.")
        finally:
            db.close()
    
    def generate_summary(self, doc_id: str) -> Dict[str, str]:
        """
        Genera un resumen del documento utilizando el LLM.
        """
        db = SessionLocal()
        
        try:
            # Verificar acceso al documento
            document = db.query(Document).filter(
                Document.id == doc_id,
                Document.tenant_id == self.tenant_id
            ).first()
            
            if not document:
                raise HTTPException(status_code=404, detail="Document not found")
            
            # Obtener texto de los chunks
            chunks = db.query(DocumentChunk).filter(
                DocumentChunk.document_id == doc_id
            ).order_by(
                DocumentChunk.chunk_index
            ).all()
            
            if not chunks:
                raise HTTPException(status_code=404, detail="Document content not found")
            
            # Unir texto de los chunks
            text = "\n\n".join([chunk.content for chunk in chunks])
            
            # Limitar longitud si es demasiado grande
            max_chars = 100000
            if len(text) > max_chars:
                text = text[:max_chars] + "..."
            
            # Generar resumen usando el servicio LLM
            summary = self.llm_service.summarize_text(text)
            
            return {"summary": summary}
            
        except HTTPException:
            raise
        except Exception as e:
            logger.exception(f"Error generating summary for document {doc_id}: {str(e)}")
            raise HTTPException(status_code=500, detail="An unexpected error occurred while generating the summary.")
        finally:
            db.close()
    
    def add_tag(self, doc_id: str, tag_name: str) -> Dict[str, str]:
        """
        Añade una etiqueta a un documento.
        """
        db = SessionLocal()
        
        try:
            document = db.query(Document).options(
                joinedload(Document.tags)
            ).filter(
                Document.id == doc_id,
                Document.tenant_id == self.tenant_id
            ).first()
            
            if not document:
                raise HTTPException(status_code=404, detail="Document not found")
            
            # Verificar si la etiqueta ya está asignada
            for tag in document.tags:
                if tag.name == tag_name:
                    return {"message": f"Tag '{tag_name}' already assigned to document"}
            
            # Buscar etiqueta existente o crear nueva
            tag = db.query(Tag).filter(
                Tag.name == tag_name,
                Tag.tenant_id == self.tenant_id
            ).first()
            
            if not tag:
                tag = Tag(name=tag_name, tenant_id=self.tenant_id)
                db.add(tag)
            
            # Asignar etiqueta al documento
            document.tags.append(tag)
            db.commit()
            
            return {"message": f"Tag '{tag_name}' added to document"}
            
        except HTTPException:
            raise
        except Exception as e:
            db.rollback()
            logger.exception(f"Error adding tag to document {doc_id}: {str(e)}")
            raise HTTPException(status_code=500, detail="An unexpected error occurred while adding the tag.")
        finally:
            db.close()
    
    def remove_tag(self, doc_id: str, tag_name: str) -> Dict[str, str]:
        """
        Elimina una etiqueta de un documento.
        """
        db = SessionLocal()
        
        try:
            document = db.query(Document).options(
                joinedload(Document.tags)
            ).filter(
                Document.id == doc_id,
                Document.tenant_id == self.tenant_id
            ).first()
            
            if not document:
                raise HTTPException(status_code=404, detail="Document not found")
            
            # Buscar etiqueta
            tag = db.query(Tag).filter(
                Tag.name == tag_name,
                Tag.tenant_id == self.tenant_id
            ).first()
            
            if not tag:
                raise HTTPException(status_code=404, detail=f"Tag '{tag_name}' not found")
            
            # Eliminar asociación
            if tag in document.tags:
                document.tags.remove(tag)
                db.commit()
                return {"message": f"Tag '{tag_name}' removed from document"}
            else:
                return {"message": f"Tag '{tag_name}' not assigned to document"}
            
        except HTTPException:
            raise
        except Exception as e:
            db.rollback()
            logger.exception(f"Error removing tag from document {doc_id}: {str(e)}")
            raise HTTPException(status_code=500, detail="An unexpected error occurred while removing the tag.")
        finally:
            db.close()
    
    def get_signed_download_url(self, doc_id: str) -> Dict[str, Any]:
        """
        Genera una URL firmada para descargar un documento.
        """
        db = SessionLocal()
        
        try:
            document = db.query(Document).filter(
                Document.id == doc_id,
                Document.tenant_id == self.tenant_id
            ).first()
            
            if not document:
                raise HTTPException(status_code=404, detail="Document not found")
            
            # Generar URL firmada
            url, expires_at = self.storage_service.generate_download_signed_url(
                object_name=document.file_path
            )
            
            return {
                "url": url,
                "expires_at": expires_at.isoformat(),
                "filename": document.filename
            }
            
        except HTTPException:
            raise
        except Exception as e:
            logger.exception(f"Error generating signed URL for document {doc_id}: {str(e)}")
            raise HTTPException(status_code=500, detail="An unexpected error occurred while generating the download URL.")
        finally:
            db.close()
    
    def get_signed_upload_url(self, filename: str, content_type: str) -> Dict[str, Any]:
        """
        Genera una URL firmada para subir un documento.
        """
        try:
            # Verificar extensión
            file_ext = os.path.splitext(filename)[1][1:].lower() if "." in filename else ""
            if not file_ext or file_ext not in settings.ALLOWED_EXTENSIONS:
                raise HTTPException(
                    status_code=400, 
                    detail=f"Tipo de archivo no permitido. Permitidos: {', '.join(settings.ALLOWED_EXTENSIONS)}"
                )
            
            # Generar ID único para el futuro documento
            doc_id = str(uuid.uuid4())
            
            # Definir ruta en el almacenamiento
            file_path = f"documents/{doc_id}/{filename}"
            
            # Generar URL firmada
            url, expires_at = self.storage_service.generate_upload_signed_url(
                object_name=file_path,
                content_type=content_type
            )
            
            return {
                "upload_url": url,
                "expires_at": expires_at.isoformat(),
                "file_path": file_path,
                "doc_id": doc_id
            }
            
        except HTTPException:
            raise
        except Exception as e:
            logger.exception(f"Error generating signed upload URL: {str(e)}")
            raise HTTPException(status_code=500, detail="An unexpected error occurred while generating the upload URL.")

    def _extract_text(self, file: BinaryIO, file_type: str) -> str:
        """
        Extrae texto de un archivo según su tipo.
        """
        try:
            file.seek(0)
            
            if file_type == 'pdf':
                return self._extract_pdf_text(file)
            elif file_type in ['docx', 'doc']:
                return self._extract_docx_text(file)
            elif file_type == 'txt':
                content = file.read()
                detected = chardet.detect(content)
                encoding = detected['encoding'] or 'utf-8'
                return content.decode(encoding, errors='ignore')
            elif file_type == 'csv':
                return self._extract_csv_text(file)
            elif file_type in ['xlsx', 'xls']:
                return self._extract_excel_text(file)
            elif file_type == 'md':
                content = file.read()
                return content.decode('utf-8', errors='ignore')
            else:
                logger.warning(f"Unsupported file type for text extraction: {file_type}")
                return ""
                
        except Exception as e:
            logger.exception(f"Error extracting text from document: {str(e)}")
            return ""
    
    def _extract_pdf_text(self, file: BinaryIO) -> str:
        """Extrae texto de un archivo PDF"""
        text = ""
        try:
            pdf_reader = PyPDF2.PdfReader(file)
            for page_num in range(len(pdf_reader.pages)):
                page = pdf_reader.pages[page_num]
                text += page.extract_text() + "\n\n"
            return text
        except Exception as e:
            logger.exception(f"Error extracting text from PDF: {str(e)}")
            return ""
    
    def _extract_docx_text(self, file: BinaryIO) -> str:
        """Extrae texto de un archivo DOCX"""
        try:
            doc = DocxDocument(file)
            text = ""
            for para in doc.paragraphs:
                text += para.text + "\n"
            return text
        except Exception as e:
            logger.exception(f"Error extracting text from DOCX: {str(e)}")
            return ""
    
    def _extract_csv_text(self, file: BinaryIO) -> str:
        """Extrae texto de un archivo CSV"""
        try:
            text = ""
            file.seek(0)
            sample = file.read(4096)
            detected = chardet.detect(sample)
            encoding = detected['encoding'] or 'utf-8'
            
            file.seek(0)
            content = file.read().decode(encoding, errors='ignore')
            file_content = io.StringIO(content)
            
            dialect = csv.Sniffer().sniff(file_content.read(1024))
            file_content.seek(0)
            
            csv_reader = csv.reader(file_content, dialect)
            for row in csv_reader:
                text += ", ".join(row) + "\n"
            return text
        except Exception as e:
            logger.exception(f"Error extracting text from CSV: {str(e)}")
            return ""
    
    def _extract_excel_text(self, file: BinaryIO) -> str:
        """Extrae texto de un archivo Excel"""
        try:
            text = ""
            workbook = openpyxl.load_workbook(file, read_only=True)
            
            for sheet_name in workbook.sheetnames:
                sheet = workbook[sheet_name]
                text += f"Sheet: {sheet_name}\n"
                
                for row in sheet.iter_rows(values_only=True):
                    row_text = ", ".join([str(cell) if cell is not None else "" for cell in row])
                    text += row_text + "\n"
                
                text += "\n"
            
            return text
        except Exception as e:
            logger.exception(f"Error extracting text from Excel: {str(e)}")
            return ""
    
    def get_documents(
        self, 
        db: "Session", # Added db: Session
        page: int = 1, 
        per_page: int = 10, 
        tags: List[str] = None, 
        date_from: str = None, 
        date_to: str = None
    ) -> Dict[str, Any]:
        """
        Obtiene lista paginada de documentos con filtros opcionales.
        
        Args:
            db: SQLAlchemy Session
            page: Número de página
            per_page: Documentos por página
            tags: Lista de etiquetas para filtrar
            date_from: Fecha inicial (formato ISO)
            date_to: Fecha final (formato ISO)
            
        Returns:
            Diccionario con documentos y metadatos de paginación
        """
        # db = next(get_db()) # Removed this line
        
        try:
            query = db.query(Document).filter(Document.tenant_id == self.tenant_id)
            
            # Aplicar filtros
            if tags:
                for tag_name in tags:
                    query = query.filter(Document.tags.any(Tag.name == tag_name))
            
            if date_from:
                try:
                    from_date = datetime.datetime.fromisoformat(date_from)
                    query = query.filter(Document.created_at >= from_date)
                except ValueError:
                    # logger.warning(f"Invalid date_from format: {date_from}")
                    raise HTTPException(status_code=400, detail="Invalid date format provided for 'date_from'. Please use ISO format.")
            
            if date_to:
                try:
                    to_date = datetime.datetime.fromisoformat(date_to)
                    query = query.filter(Document.created_at <= to_date)
                except ValueError:
                    # logger.warning(f"Invalid date_to format: {date_to}")
                    raise HTTPException(status_code=400, detail="Invalid date format provided for 'date_to'. Please use ISO format.")
            
            # Contar total
            total = query.count()
            
            # Aplicar paginación
            offset = (page - 1) * per_page
            documents = query.order_by(Document.created_at.desc()).offset(offset).limit(per_page).all()
            
            # Convertir a diccionarios
            documents_dict = []
            for doc in documents:
                doc_dict = {
                    "id": str(doc.id),
                    "title": doc.title,
                    "description": doc.description,
                    "filename": doc.filename,
                    "file_path": doc.file_path,
                    "file_type": doc.file_type,
                    "file_size": doc.file_size,
                    "indexed": doc.indexed,
                    "created_at": doc.created_at.isoformat(),
                    "updated_at": doc.updated_at.isoformat(),
                    "tags": [tag.name for tag in doc.tags]
                }
                documents_dict.append(doc_dict)
            
            return {
                "documents": documents_dict,
                "pagination": {
                    "page": page,
                    "per_page": per_page,
                    "total": total,
                    "pages": (total + per_page - 1) // per_page
                }
            }
            
        except Exception as e:
            logger.exception(f"Error getting documents: {str(e)}")
            raise HTTPException(status_code=500, detail="An unexpected error occurred while retrieving documents.")
    
    def get_document(self, db: "Session", doc_id: str) -> Dict[str, Any]: # Added db: Session
        """
        Obtiene información detallada de un documento.
        
        Args:
            db: SQLAlchemy Session
            doc_id: ID del documento
            
        Returns:
            Diccionario con información del documento
        """
        # db = next(get_db()) # Removed this line
        
        try:
            document = db.query(Document).filter(
                Document.id == doc_id,
                Document.tenant_id == self.tenant_id
            ).first()
            
            if not document:
                raise HTTPException(status_code=404, detail="Document not found")
            
            # Obtener primeros chunks para vista previa
            chunks = db.query(DocumentChunk).filter(
                DocumentChunk.document_id == doc_id
            ).order_by(
                DocumentChunk.chunk_index
            ).limit(3).all()
            
            chunks_dict = []
            for chunk in chunks:
                chunks_dict.append({
                    "id": chunk.id,
                    "document_id": str(chunk.document_id),
                    "chunk_index": chunk.chunk_index,
                    "content": chunk.content
                })
            
            result = {
                "id": str(document.id),
                "title": document.title,
                "description": document.description,
                "filename": document.filename,
                "file_path": document.file_path,
                "file_type": document.file_type,
                "file_size": document.file_size,
                "indexed": document.indexed,
                "created_at": document.created_at.isoformat(),
                "updated_at": document.updated_at.isoformat(),
                "tags": [tag.name for tag in document.tags],
                "preview_chunks": chunks_dict
            }
            
            return result
            
        except HTTPException:
            raise
        except Exception as e:
            logger.exception(f"Error getting document {doc_id}: {str(e)}")
            raise HTTPException(status_code=500, detail="An unexpected error occurred while retrieving the document.")
    
    def delete_document(self, db: "Session", doc_id: str) -> Dict[str, Any]: # Added db: Session
        """
        Elimina un documento y todos sus datos asociados.
        
        Args:
            db: SQLAlchemy Session
            doc_id: ID del documento
            
        Returns:
            Mensaje de confirmación
        """
        # db = next(get_db()) # Removed this line
        
        try:
            document = db.query(Document).filter(
                Document.id == doc_id,
                Document.tenant_id == self.tenant_id
            ).first()
            
            if not document:
                raise HTTPException(status_code=404, detail="Document not found")
            
            # Eliminar archivo del almacenamiento
            self.storage_service.delete_file(document.file_path)
            
            # Eliminar del vector store
            self.embedding_service.delete_document(doc_id)
            
            # Eliminar de la base de datos
            db.delete(document)
            db.commit()
            
            return {"message": f"Document {doc_id} deleted successfully"}
            
        except HTTPException:
            raise
        except Exception as e:
            db.rollback()
            logger.exception(f"Error deleting document {doc_id}: {str(e)}")
            raise HTTPException(status_code=500, detail="An unexpected error occurred while deleting the document.")
    
    def generate_summary(self, db: "Session", doc_id: str) -> Dict[str, str]: # Added db: Session
        """
        Genera un resumen del documento utilizando el LLM.
        
        Args:
            db: SQLAlchemy Session
            doc_id: ID del documento
            
        Returns:
            Texto del resumen
        """
        # db = next(get_db()) # Removed this line
        
        try:
            # Verificar acceso al documento
            document = db.query(Document).filter(
                Document.id == doc_id,
                Document.tenant_id == self.tenant_id
            ).first()
            
            if not document:
                raise HTTPException(status_code=404, detail="Document not found")
            
            # Obtener texto de los chunks
            chunks = db.query(DocumentChunk).filter(
                DocumentChunk.document_id == doc_id
            ).order_by(
                DocumentChunk.chunk_index
            ).all()
            
            if not chunks:
                raise HTTPException(status_code=404, detail="Document content not found")
            
            # Unir texto de los chunks
            text = "\n\n".join([chunk.content for chunk in chunks])
            
            # Limitar longitud si es demasiado grande
            max_chars = 100000  # Ajustar según limitaciones del LLM
            if len(text) > max_chars:
                text = text[:max_chars] + "..."
            
            # Generar resumen usando el servicio LLM
            summary = self.llm_service.summarize_text(text)
            
            return {"summary": summary}
            
        except HTTPException:
            raise
        except Exception as e:
            logger.exception(f"Error generating summary for document {doc_id}: {str(e)}")
            raise HTTPException(status_code=500, detail="An unexpected error occurred while generating the summary.")
    
    def add_tag(self, db: "Session", doc_id: str, tag_name: str) -> Dict[str, str]: # Added db: Session
        """
        Añade una etiqueta a un documento.
        
        Args:
            db: SQLAlchemy Session
            doc_id: ID del documento
            tag_name: Nombre de la etiqueta
            
        Returns:
            Mensaje de confirmación
        """
        # db = next(get_db()) # Removed this line
        
        try:
            document = db.query(Document).filter(
                Document.id == doc_id,
                Document.tenant_id == self.tenant_id
            ).first()
            
            if not document:
                raise HTTPException(status_code=404, detail="Document not found")
            
            # Verificar si la etiqueta ya está asignada
            for tag in document.tags:
                if tag.name == tag_name:
                    return {"message": f"Tag '{tag_name}' already assigned to document"}
            
            # Buscar etiqueta existente o crear nueva
            tag = db.query(Tag).filter(
                Tag.name == tag_name,
                Tag.tenant_id == self.tenant_id
            ).first()
            
            if not tag:
                tag = Tag(name=tag_name, tenant_id=self.tenant_id)
                db.add(tag)
            
            # Asignar etiqueta al documento
            document.tags.append(tag)
            db.commit()
            
            return {"message": f"Tag '{tag_name}' added to document"}
            
        except HTTPException:
            raise
        except Exception as e:
            db.rollback()
            logger.exception(f"Error adding tag to document {doc_id}: {str(e)}")
            raise HTTPException(status_code=500, detail="An unexpected error occurred while adding the tag.")
    
    def remove_tag(self, db: "Session", doc_id: str, tag_name: str) -> Dict[str, str]: # Added db: Session
        """
        Elimina una etiqueta de un documento.
        
        Args:
            db: SQLAlchemy Session
            doc_id: ID del documento
            tag_name: Nombre de la etiqueta
            
        Returns:
            Mensaje de confirmación
        """
        # db = next(get_db()) # Removed this line
        
        try:
            document = db.query(Document).filter(
                Document.id == doc_id,
                Document.tenant_id == self.tenant_id
            ).first()
            
            if not document:
                raise HTTPException(status_code=404, detail="Document not found")
            
            # Buscar etiqueta
            tag = db.query(Tag).filter(
                Tag.name == tag_name,
                Tag.tenant_id == self.tenant_id
            ).first()
            
            if not tag:
                raise HTTPException(status_code=404, detail=f"Tag '{tag_name}' not found")
            
            # Eliminar asociación
            if tag in document.tags:
                document.tags.remove(tag)
                db.commit()
                return {"message": f"Tag '{tag_name}' removed from document"}
            else:
                return {"message": f"Tag '{tag_name}' not assigned to document"}
            
        except HTTPException:
            raise
        except Exception as e:
            db.rollback()
            logger.exception(f"Error removing tag from document {doc_id}: {str(e)}")
            raise HTTPException(status_code=500, detail="An unexpected error occurred while removing the tag.")
    
    def get_signed_download_url(self, db: "Session", doc_id: str) -> Dict[str, Any]: # Added db: Session
        """
        Genera una URL firmada para descargar un documento.
        
        Args:
            db: SQLAlchemy Session
            doc_id: ID del documento
            
        Returns:
            URL firmada y fecha de expiración
        """
        # db = next(get_db()) # Removed this line
        
        try:
            document = db.query(Document).filter(
                Document.id == doc_id,
                Document.tenant_id == self.tenant_id
            ).first()
            
            if not document:
                raise HTTPException(status_code=404, detail="Document not found")
            
            # Generar URL firmada
            url, expires_at = self.storage_service.generate_download_signed_url(
                object_name=document.file_path
            )
            
            return {
                "url": url,
                "expires_at": expires_at.isoformat(),
                "filename": document.filename
            }
            
        except HTTPException:
            raise
        except Exception as e:
            logger.exception(f"Error generating signed URL for document {doc_id}: {str(e)}")
            raise HTTPException(status_code=500, detail="An unexpected error occurred while generating the download URL.")
    
    def get_signed_upload_url(self, filename: str, content_type: str) -> Dict[str, Any]:
        """
        Genera una URL firmada para subir un documento.
        
        Args:
            filename: Nombre del archivo
            content_type: Tipo de contenido MIME
            
        Returns:
            URL firmada, fecha de expiración y path de destino
        """
        try:
            # Verificar extensión
            file_ext = os.path.splitext(filename)[1][1:].lower() if "." in filename else ""
            if not file_ext or file_ext not in settings.ALLOWED_EXTENSIONS:
                raise HTTPException(
                    status_code=400, 
                    detail=f"Tipo de archivo no permitido. Permitidos: {', '.join(settings.ALLOWED_EXTENSIONS)}"
                )
            
            # Generar ID único para el futuro documento
            doc_id = str(uuid.uuid4())
            
            # Definir ruta en el almacenamiento
            file_path = f"documents/{doc_id}/{filename}"
            
            # Generar URL firmada
            url, expires_at = self.storage_service.generate_upload_signed_url(
                object_name=file_path,
                content_type=content_type
            )
            
            return {
                "upload_url": url,
                "expires_at": expires_at.isoformat(),
                "file_path": file_path,
                "doc_id": doc_id
            }
            
        except HTTPException:
            raise
        except Exception as e:
            logger.exception(f"Error generating signed upload URL: {str(e)}")
            raise HTTPException(status_code=500, detail="An unexpected error occurred while generating the upload URL.")