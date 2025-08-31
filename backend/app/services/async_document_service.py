"""
Async version of DocumentService for use with AsyncSession
"""
import logging
import os
import uuid
import datetime
import io
import asyncio
from typing import List, Dict, Any, Optional, BinaryIO, Union
from fastapi import UploadFile, HTTPException
import PyPDF2
from docx import Document as DocxDocument
import csv
import openpyxl
import chardet
from sqlalchemy.ext.asyncio import AsyncSession
from sqlalchemy import select, func, and_, or_
from sqlalchemy.orm import selectinload, joinedload

from app.core.config import settings
from app.db.models import Document, Tag, Tenant
from app.db.async_database import AsyncSessionLocal
from app.schemas.enums import IndexingStatus
from app.services.async_storage_factory import AsyncStorageServiceFactory
from app.services.embedding_service import EmbeddingService
from app.services.vector_service_direct import VectorServiceDirect
from app.services.llm_service import LLMService

logger = logging.getLogger(__name__)


class AsyncDocumentService:
    """Async version of Document Service"""
    
    def __init__(self, tenant_id: str = None, user_id: str = None):
        """
        Initialize async document service
        Note: The async init pattern requires using a factory method
        """
        self.tenant_id = tenant_id
        self.user_id = user_id
        self.storage_service = None
        self.embedding_service = None
        self.vector_service = None
        self.llm_service = None
        self._initialized = False
    
    @classmethod
    async def create(cls, tenant_id: str = None, user_id: str = None, db: AsyncSession = None):
        """
        Factory method to create and initialize AsyncDocumentService
        """
        service = cls(tenant_id, user_id)
        await service._initialize(db)
        return service
    
    async def _initialize(self, db: AsyncSession = None):
        """Initialize the service with async operations"""
        # If tenant_id is None or "default", get the real UUID
        if not self.tenant_id or self.tenant_id == settings.DEFAULT_TENANT:
            if db:
                # Use provided session
                stmt = select(Tenant).filter(Tenant.name == settings.DEFAULT_TENANT)
                result = await db.execute(stmt)
                default_tenant = result.scalar_one_or_none()
                
                if default_tenant:
                    self.tenant_id = str(default_tenant.id)
                else:
                    raise ValueError(f"Default tenant '{settings.DEFAULT_TENANT}' not found in database")
                
                # Create storage service using async factory
                self.storage_service = await AsyncStorageServiceFactory.create_storage_service(
                    self.tenant_id, self.user_id, db
                )
            else:
                # Create new session only if not provided
                async with AsyncSessionLocal() as new_db:
                    stmt = select(Tenant).filter(Tenant.name == settings.DEFAULT_TENANT)
                    result = await new_db.execute(stmt)
                    default_tenant = result.scalar_one_or_none()
                    
                    if default_tenant:
                        self.tenant_id = str(default_tenant.id)
                    else:
                        raise ValueError(f"Default tenant '{settings.DEFAULT_TENANT}' not found in database")
                    
                    # Create storage service using async factory
                    self.storage_service = await AsyncStorageServiceFactory.create_storage_service(
                        self.tenant_id, self.user_id, new_db
                    )
        else:
            if db:
                self.storage_service = await AsyncStorageServiceFactory.create_storage_service(
                    self.tenant_id, self.user_id, db
                )
            else:
                async with AsyncSessionLocal() as new_db:
                    self.storage_service = await AsyncStorageServiceFactory.create_storage_service(
                        self.tenant_id, self.user_id, new_db
                    )
        
        self.embedding_service = EmbeddingService(self.tenant_id)
        self.vector_service = VectorServiceDirect(self.tenant_id, self.user_id)
        self.llm_service = LLMService()
        self._initialized = True
    
    async def _validate_file(self, file: UploadFile, filename: str) -> tuple[str, bytes, int]:
        """Validates the uploaded file"""
        if not file:
            raise HTTPException(status_code=400, detail="Archivo no proporcionado")

        contents = await file.read()
        await file.seek(0) 

        file_size = len(contents)
        file_ext = os.path.splitext(filename)[1][1:].lower() if "." in filename else ""

        if not file_ext or file_ext not in settings.ALLOWED_EXTENSIONS:
            raise HTTPException(
                status_code=400,
                detail=f"Tipo de archivo no permitido. Permitidos: {', '.join(settings.ALLOWED_EXTENSIONS)}"
            )
        
        if file_size == 0:
             raise HTTPException(status_code=400, detail="El archivo está vacío.")

        if file_size > settings.MAX_UPLOAD_SIZE:
            raise HTTPException(
                status_code=400,
                detail=f"Tamaño de archivo excede el límite de {settings.MAX_UPLOAD_SIZE // (1024*1024)}MB"
            )
        
        return file_ext, contents, file_size
    
    async def get_documents(
        self, 
        db: AsyncSession,
        page: int = 1,
        per_page: int = 10,
        search: Optional[str] = None,
        tags: Optional[List[str]] = None,
        date_from: Optional[str] = None,
        date_to: Optional[str] = None,
        category: Optional[str] = None
    ) -> Dict[str, Any]:
        """Get paginated list of documents with filters"""
        try:
            # Base query
            query = select(Document).filter(
                Document.tenant_id == self.tenant_id
            ).options(
                selectinload(Document.tags),
                selectinload(Document.creator)
            )
            
            # Apply filters
            if search:
                query = query.filter(
                    or_(
                        Document.title.ilike(f"%{search}%"),
                        Document.description.ilike(f"%{search}%"),
                        Document.content.ilike(f"%{search}%")
                    )
                )
            
            if category:
                query = query.filter(Document.category == category)
            
            if tags:
                # Join with tags
                query = query.join(Document.tags).filter(
                    Tag.name.in_(tags)
                )
            
            if date_from:
                date_from_obj = datetime.datetime.fromisoformat(date_from)
                query = query.filter(Document.created_at >= date_from_obj)
            
            if date_to:
                date_to_obj = datetime.datetime.fromisoformat(date_to)
                query = query.filter(Document.created_at <= date_to_obj)
            
            # Count total
            count_query = select(func.count()).select_from(query.subquery())
            total_result = await db.execute(count_query)
            total = total_result.scalar()
            
            # Apply pagination
            offset = (page - 1) * per_page
            query = query.offset(offset).limit(per_page).order_by(Document.created_at.desc())
            
            # Execute query
            result = await db.execute(query)
            documents = result.scalars().all()
            
            # Convert to dict
            items = []
            for doc in documents:
                doc_dict = {
                    "id": str(doc.id),
                    "title": doc.title,
                    "description": doc.description,
                    "filename": doc.filename,
                    "file_type": doc.file_type,
                    "file_size": doc.file_size,
                    "mime_type": doc.mime_type,
                    "indexed": self._get_indexed_status_string(doc.indexed),
                    "category": doc.category if hasattr(doc, 'category') else None,
                    "created_at": doc.created_at.isoformat() if doc.created_at else None,
                    "updated_at": doc.updated_at.isoformat() if doc.updated_at else None,
                    "tags": [{"id": str(tag.id), "name": tag.name} for tag in doc.tags],
                    "created_by": {
                        "id": str(doc.creator.id),
                        "email": doc.creator.email,
                        "full_name": doc.creator.full_name
                    } if doc.creator else None
                }
                items.append(doc_dict)
            
            return {
                "items": items,
                "total": total,
                "page": page,
                "per_page": per_page,
                "pages": (total + per_page - 1) // per_page
            }
            
        except Exception as e:
            logger.error(f"Error getting documents: {e}")
            raise HTTPException(status_code=500, detail=str(e))
    
    async def upload_document(
        self,
        db: AsyncSession,
        file: UploadFile,
        title: str,
        description: Optional[str] = None,
        tags: Optional[List[str]] = None,
        category: Optional[str] = None
    ) -> Document:
        """Upload a new document"""
        try:
            # Ensure service is initialized
            if not self._initialized:
                await self._initialize(db)
            # Validate file
            file_ext, contents, file_size = await self._validate_file(file, file.filename)
            
            # Generate unique filename
            file_id = str(uuid.uuid4())
            stored_filename = f"{file_id}.{file_ext}"
            
            # Create document record
            doc = Document(
                id=uuid.uuid4(),
                title=title,
                description=description,
                filename=file.filename,
                file_path=stored_filename,
                file_type=file_ext,
                file_size=file_size,
                mime_type=file.content_type,
                category=category,
                tenant_id=self.tenant_id,
                created_by=self.user_id,
                indexed=IndexingStatus.PROCESSING  # Set initial status
            )
            
            # Upload to storage
            upload_success = await self.storage_service.upload_file(
                file=io.BytesIO(contents),
                object_name=stored_filename,
                metadata={"content_type": file.content_type}
            )
            
            if not upload_success:
                raise HTTPException(status_code=500, detail="Error uploading file")
            
            # Add to database
            db.add(doc)
            await db.flush()
            
            # Handle tags
            if tags:
                for tag_name in tags:
                    # Check if tag exists
                    stmt = select(Tag).filter(
                        Tag.name == tag_name,
                        Tag.tenant_id == self.tenant_id
                    )
                    result = await db.execute(stmt)
                    tag = result.scalar_one_or_none()
                    
                    if not tag:
                        tag = Tag(
                            id=uuid.uuid4(),
                            name=tag_name,
                            tenant_id=self.tenant_id
                        )
                        db.add(tag)
                    
                    doc.tags.append(tag)
            
            await db.commit()
            
            # Reload the document with proper eager loading for tags
            stmt = select(Document).filter(
                Document.id == doc.id
            ).options(
                selectinload(Document.tags),
                selectinload(Document.creator)
            )
            result = await db.execute(stmt)
            doc = result.scalar_one()
            
            # Extract text and index asynchronously
            # Pass document info to avoid needing new DB session
            doc_info = {
                "id": str(doc.id),
                "filename": doc.filename,
                "tenant_id": self.tenant_id,
                "user_id": self.user_id
            }
            task = asyncio.create_task(self._process_document_async(doc_info, contents, file_ext))
            # Add error handler for the background task
            task.add_done_callback(lambda t: logger.error(f"Background processing failed: {t.exception()}") if t.exception() else None)
            
            return doc
            
        except Exception as e:
            await db.rollback()
            logger.error(f"Error uploading document: {e}")
            raise HTTPException(status_code=500, detail=str(e))
    
    async def _process_document_async(self, doc_info: dict, contents: bytes, file_ext: str):
        """Process document in background"""
        doc_id = doc_info["id"]
        try:
            logger.info(f"Starting async processing for document {doc_id}, file type: {file_ext}")
            
            # Extract text
            text = await self._extract_text_async(contents, file_ext)
            logger.info(f"Text extraction completed for {doc_id}, text length: {len(text) if text else 0}")
            
            if text:
                try:
                    # Limit text length for embedding generation to avoid timeouts
                    # With all-minilm, we can process text faster but let's be conservative
                    max_text_length = 30000  # Limit to ~30k characters for faster processing
                    if len(text) > max_text_length:
                        logger.warning(f"Text too long ({len(text)} chars), truncating to {max_text_length} for embeddings")
                        text_for_embedding = text[:max_text_length]
                    else:
                        text_for_embedding = text
                    
                    # Generate embeddings
                    logger.info(f"Generating embeddings for document {doc_id} (text length: {len(text_for_embedding)})")
                    embeddings = await self.embedding_service.generate_embeddings(text_for_embedding)
                    logger.info(f"Generated {len(embeddings)} embeddings for document {doc_id}")
                except Exception as e:
                    logger.error(f"Failed to generate embeddings for {doc_id}: {e}")
                    raise Exception(f"Embedding generation failed: {str(e)}")
                
                try:
                    # Get full document info for metadata before storing in vector DB
                    async with AsyncSessionLocal() as db:
                        stmt = select(Document).filter(Document.id == doc_id)
                        result = await db.execute(stmt)
                        doc = result.scalar_one_or_none()
                        
                        if not doc:
                            raise Exception("Document not found in database")
                    
                    # Build comprehensive metadata for vector search
                    metadata = {
                        "file_type": file_ext,
                        "tenant_id": self.tenant_id,
                        "filename": doc.filename,
                        "title": doc.title or doc.filename,
                        "description": doc.description,
                        "created_at": doc.created_at.isoformat() if doc.created_at else None,
                        "updated_at": doc.updated_at.isoformat() if doc.updated_at else None,
                        "file_size": doc.file_size,
                        "mime_type": doc.mime_type,
                        "category": doc.category
                    }
                    
                    # Store in vector DB (use the same text that was used for embeddings)
                    logger.info(f"Storing document {doc_id} in vector database with enhanced metadata")
                    success = await self.vector_service.add_document(
                        doc_id=doc_id,
                        text=text_for_embedding,  # Use the same text that was embedded
                        metadata=metadata
                    )
                    
                    if not success:
                        raise Exception("Failed to store document in vector database")
                    logger.info(f"Successfully stored document {doc_id} in vector database")
                except Exception as e:
                    logger.error(f"Failed to store in vector DB for {doc_id}: {e}")
                    raise Exception(f"Vector storage failed: {str(e)}")
                
                # Update document status and perform routing in single session
                async with AsyncSessionLocal() as db:
                    stmt = select(Document).filter(Document.id == doc_id)
                    result = await db.execute(stmt)
                    doc = result.scalar_one_or_none()
                    
                    if doc:
                        doc.indexed = IndexingStatus.INDEXED
                        # Generate summary instead of storing first 1000 chars
                        summary = await self._generate_document_summary(text, doc_info["filename"])
                        doc.content = summary[:1000]  # Store summary (max 1000 chars)
                        await db.commit()
                        logger.info(f"Document {doc_id} marked as INDEXED with summary")
                        
                        # Perform routing analysis in the same session
                        await self._perform_routing_analysis(db, doc_info, text, file_ext)
                
                # Queue document for auto-categorization
                try:
                    from app.services.queue_service import queue_service
                    await queue_service.enqueue_document_categorization(
                        document_id=doc_id,
                        tenant_id=self.tenant_id,
                        user_id=self.user_id,
                        priority="default"
                    )
                    logger.info(f"Document {doc_id} queued for categorization")
                except Exception as e:
                    logger.warning(f"Failed to queue categorization for {doc_id}: {e}")
                    # Don't fail the whole process if categorization queueing fails
                
                # Routing is now handled in the same DB session above to avoid greenlet errors
            else:
                logger.warning(f"No text extracted from document {doc_id}")
                raise Exception("No text could be extracted from the document")
            
        except Exception as e:
            logger.error(f"Error processing document {doc_id}: {e}", exc_info=True)
            # Update error status
            async with AsyncSessionLocal() as db:
                stmt = select(Document).filter(Document.id == doc_id)
                result = await db.execute(stmt)
                doc = result.scalar_one_or_none()
                
                if doc:
                    doc.indexed = IndexingStatus.INDEXING_ERROR
                    doc.indexing_error = str(e)
                    await db.commit()
                    logger.info(f"Document {doc_id} marked as INDEXING_ERROR: {str(e)}")
    
    async def _extract_text_async(self, contents: bytes, file_ext: str) -> Optional[str]:
        """Extract text from document asynchronously"""
        # This is a simplified version - in production you'd want proper async extraction
        return await asyncio.to_thread(self._extract_text_sync, contents, file_ext)
    
    def _extract_text_sync(self, contents: bytes, file_ext: str) -> Optional[str]:
        """Synchronous text extraction with OCR fallback for scanned PDFs"""
        try:
            if file_ext == "pdf":
                # First try PyPDF2 for text-based PDFs
                pdf_reader = PyPDF2.PdfReader(io.BytesIO(contents))
                text = ""
                for page in pdf_reader.pages:
                    text += page.extract_text() + "\n"
                
                # If no text extracted (likely scanned PDF), try OCR
                if not text.strip():
                    logger.info("No text extracted with PyPDF2, attempting OCR for scanned PDF")
                    try:
                        import fitz  # PyMuPDF
                        from PIL import Image
                        import pytesseract
                        
                        # Open PDF with PyMuPDF for better image handling
                        doc = fitz.open(stream=contents, filetype="pdf")
                        ocr_text = ""
                        
                        for page_num in range(len(doc)):
                            page = doc.load_page(page_num)
                            # Convert page to image
                            mat = fitz.Matrix(2, 2)  # 2x zoom for better OCR
                            pix = page.get_pixmap(matrix=mat)
                            img_data = pix.tobytes("png")
                            
                            # OCR the image
                            image = Image.open(io.BytesIO(img_data))
                            page_text = pytesseract.image_to_string(image, lang='spa+eng')  # Spanish + English
                            ocr_text += page_text + "\n"
                            logger.info(f"OCR extracted {len(page_text)} chars from page {page_num + 1}")
                        
                        doc.close()
                        text = ocr_text
                        logger.info(f"OCR completed: extracted {len(text)} total characters")
                        
                    except ImportError:
                        logger.warning("OCR libraries not available (fitz, PIL, pytesseract). Install: pip install PyMuPDF Pillow pytesseract")
                        return ""
                    except Exception as ocr_error:
                        logger.error(f"OCR failed: {ocr_error}")
                        return ""
                
                return text
            
            elif file_ext == "txt":
                encoding = chardet.detect(contents)['encoding'] or 'utf-8'
                return contents.decode(encoding)
            
            elif file_ext == "docx":
                doc = DocxDocument(io.BytesIO(contents))
                return "\n".join([paragraph.text for paragraph in doc.paragraphs])
            
            else:
                return None
                
        except Exception as e:
            logger.error(f"Error extracting text: {e}")
            return None
    
    async def get_document(self, db: AsyncSession, doc_id: str) -> Document:
        """Get single document by ID"""
        stmt = select(Document).filter(
            Document.id == doc_id,
            Document.tenant_id == self.tenant_id
        ).options(
            selectinload(Document.tags),
            selectinload(Document.creator)
        )
        
        result = await db.execute(stmt)
        doc = result.scalar_one_or_none()
        
        if not doc:
            raise HTTPException(status_code=404, detail="Document not found")
        
        return doc
    
    async def delete_document(self, db: AsyncSession, doc_id: str) -> Dict[str, Any]:
        """Delete a document"""
        doc = await self.get_document(db, doc_id)
        
        # Mark document as deleted by removing it from database
        # Note: Consider implementing soft delete with is_active field if needed
        
        # Delete from vector DB
        try:
            await self.vector_service.delete_document(str(doc.id))
        except Exception as e:
            logger.error(f"Error deleting from vector DB: {e}")
        
        # Delete from storage
        try:
            await self.storage_service.delete_file(doc.file_path)
        except Exception as e:
            logger.error(f"Error deleting from storage: {e}")
        
        # Delete from database
        await db.delete(doc)
        await db.commit()
        
        return {"success": True, "message": "Document deleted successfully"}
    
    async def _auto_categorize_document(self, doc_id: str, text_content: str):
        """Auto-categorize document using the LangChain service"""
        try:
            import httpx
            from app.core.config import settings
            
            async with AsyncSessionLocal() as db:
                # Get document info
                stmt = select(Document).filter(Document.id == doc_id)
                result = await db.execute(stmt)
                doc = result.scalar_one_or_none()
                
                if not doc:
                    return
                
                # First try LangChain service for simple categorization
                try:
                    async with httpx.AsyncClient() as client:
                        # Use LangChain service for document analysis
                        response = await client.post(
                            f"{settings.LANGCHAIN_SERVICE_URL}/api/v1/chat/completions",
                            json={
                                "messages": [{
                                    "role": "system",
                                    "content": """You are a document categorization expert. Analyze the document and categorize it into one of these categories:
                                    - contract: Legal contracts, agreements, terms
                                    - invoice: Invoices, bills, receipts
                                    - report: Reports, analysis, research documents
                                    - legal: Legal documents, policies, regulations
                                    - financial: Financial statements, budgets, accounting
                                    - technical: Technical documentation, manuals, specifications
                                    - correspondence: Letters, emails, memos
                                    - presentation: Slides, presentations
                                    - general: Other documents
                                    
                                    Respond with ONLY the category name, nothing else."""
                                }, {
                                    "role": "user", 
                                    "content": f"Document name: {doc.filename}\nContent preview: {text_content[:1000]}"
                                }],
                                "model": settings.OLLAMA_MODEL,
                                "max_tokens": 50,
                                "temperature": 0.1
                            },
                            headers={
                                "X-API-Key": settings.MICROSERVICES_API_KEY,
                                "Content-Type": "application/json"
                            },
                            timeout=30.0
                        )
                        
                        if response.status_code == 200:
                            result = response.json()
                            category = result.get("choices", [{}])[0].get("message", {}).get("content", "").strip().lower()
                            
                            # Validate category
                            valid_categories = ["contract", "invoice", "report", "legal", "financial", 
                                              "technical", "correspondence", "presentation", "general"]
                            if category not in valid_categories:
                                category = "general"
                            
                            # Update document category
                            doc.category = category
                            doc.document_metadata = doc.document_metadata or {}
                            doc.document_metadata["auto_categorization"] = {
                                "timestamp": datetime.datetime.now(datetime.timezone.utc).isoformat(),
                                "method": "langchain",
                                "model": settings.OLLAMA_MODEL
                            }
                            await db.commit()
                            
                            logger.info(f"Document {doc_id} auto-categorized as '{category}'")
                        else:
                            logger.warning(f"Failed to auto-categorize document {doc_id}: HTTP {response.status_code}")
                            
                except Exception as e:
                    logger.warning(f"LangChain categorization failed for {doc_id}: {e}")
                    # Fallback to simple rule-based categorization
                    category = self._simple_categorize(doc.filename, text_content)
                    doc.category = category
                    doc.document_metadata = doc.document_metadata or {}
                    doc.document_metadata["auto_categorization"] = {
                        "timestamp": datetime.datetime.now(datetime.timezone.utc).isoformat(),
                        "method": "rule_based"
                    }
                    await db.commit()
                    logger.info(f"Document {doc_id} categorized as '{category}' using rules")
                        
        except Exception as e:
            logger.error(f"Error auto-categorizing document {doc_id}: {e}")
            # Don't fail the whole process if categorization fails
    
    def _simple_categorize(self, filename: str, content: str) -> str:
        """Simple rule-based categorization as fallback"""
        filename_lower = filename.lower()
        content_lower = content.lower()[:1000]  # Check first 1000 chars
        
        # Check filename and content for patterns
        if any(word in filename_lower for word in ["contract", "agreement", "terms"]):
            return "contract"
        elif any(word in filename_lower for word in ["invoice", "bill", "receipt"]):
            return "invoice"
        elif any(word in filename_lower for word in ["report", "analysis"]):
            return "report"
        elif any(word in content_lower for word in ["whereas", "agreement", "party", "shall"]):
            return "contract"
        elif any(word in content_lower for word in ["invoice", "total", "payment due", "bill to"]):
            return "invoice"
        elif any(word in content_lower for word in ["executive summary", "findings", "conclusion"]):
            return "report"
        elif filename_lower.endswith((".pptx", ".ppt")):
            return "presentation"
        else:
            return "general"
    
    def _get_indexed_status_string(self, indexed_value: int) -> str:
        """Convert indexed integer value to string representation"""
        if indexed_value == IndexingStatus.INDEXED:
            return "INDEXED"
        elif indexed_value == IndexingStatus.PROCESSING:
            return "PROCESSING"
        elif indexed_value == IndexingStatus.INDEXING_ERROR:
            return "INDEXING_ERROR"
        else:  # NOT_INDEXED or unknown
            return "NOT_INDEXED"
    
    async def generate_summary(self, db: AsyncSession, doc_id: str) -> Dict[str, str]:
        """Generate document summary using LLM"""
        try:
            # Get document
            document = await self.get_document(db, doc_id)
            if not document:
                raise HTTPException(status_code=404, detail="Document not found")
            
            # Check if document has text content
            if not document.text_content:
                raise HTTPException(status_code=400, detail="Document has no text content to summarize")
            
            # TODO: Implement actual LLM summary generation
            # For now, return a simple summary
            text_preview = document.text_content[:500] if document.text_content else ""
            word_count = len(document.text_content.split()) if document.text_content else 0
            
            return {
                "summary": f"This document contains {word_count} words. Preview: {text_preview}...",
                "status": "generated"
            }
            
        except HTTPException:
            raise
        except Exception as e:
            logger.error(f"Error generating summary for document {doc_id}: {str(e)}")
            raise HTTPException(status_code=500, detail="Error generating summary")
    
    async def add_tag(self, db: AsyncSession, doc_id: str, tag_name: str) -> Dict[str, Any]:
        """Add a tag to a document"""
        try:
            # Get document
            stmt = select(Document).filter(
                Document.id == doc_id,
                Document.tenant_id == self.tenant_id
            ).options(selectinload(Document.tags))
            
            result = await db.execute(stmt)
            document = result.scalar_one_or_none()
            
            if not document:
                raise HTTPException(status_code=404, detail="Document not found")
            
            # Check if tag already exists
            tag_stmt = select(Tag).filter(
                Tag.name == tag_name,
                Tag.tenant_id == self.tenant_id
            )
            tag_result = await db.execute(tag_stmt)
            tag = tag_result.scalar_one_or_none()
            
            if not tag:
                # Create new tag
                tag = Tag(name=tag_name, tenant_id=self.tenant_id)
                db.add(tag)
                await db.flush()
            
            # Add tag to document if not already added
            if tag not in document.tags:
                document.tags.append(tag)
                await db.commit()
            
            return {
                "message": f"Tag '{tag_name}' added successfully",
                "document_id": str(document.id),
                "tags": [t.name for t in document.tags]
            }
            
        except HTTPException:
            raise
        except Exception as e:
            await db.rollback()
            logger.error(f"Error adding tag to document {doc_id}: {str(e)}")
            raise HTTPException(status_code=500, detail="Error adding tag")
    
    async def remove_tag(self, db: AsyncSession, doc_id: str, tag_name: str) -> Dict[str, Any]:
        """Remove a tag from a document"""
        try:
            # Get document
            stmt = select(Document).filter(
                Document.id == doc_id,
                Document.tenant_id == self.tenant_id
            ).options(selectinload(Document.tags))
            
            result = await db.execute(stmt)
            document = result.scalar_one_or_none()
            
            if not document:
                raise HTTPException(status_code=404, detail="Document not found")
            
            # Find and remove the tag
            tag_to_remove = None
            for tag in document.tags:
                if tag.name == tag_name:
                    tag_to_remove = tag
                    break
            
            if tag_to_remove:
                document.tags.remove(tag_to_remove)
                await db.commit()
                
                return {
                    "message": f"Tag '{tag_name}' removed successfully",
                    "document_id": str(document.id),
                    "tags": [t.name for t in document.tags]
                }
            else:
                raise HTTPException(status_code=404, detail=f"Tag '{tag_name}' not found on document")
            
        except HTTPException:
            raise
        except Exception as e:
            await db.rollback()
            logger.error(f"Error removing tag from document {doc_id}: {str(e)}")
            raise HTTPException(status_code=500, detail="Error removing tag")
    
    async def _generate_document_summary(self, text: str, filename: str) -> str:
        """Generate a concise summary of the document"""
        try:
            import httpx
            from app.core.config import settings
            
            # Prepare text for summarization (limit to reasonable size)
            text_for_summary = text[:5000] if len(text) > 5000 else text
            
            # Try to generate summary using LangChain service
            try:
                async with httpx.AsyncClient(timeout=30.0) as client:
                    response = await client.post(
                        f"{settings.LANGCHAIN_SERVICE_URL}/api/v1/chat/completions",
                        json={
                            "messages": [{
                                "role": "system",
                                "content": "You are a document summarizer. Create a concise summary of the document in 2-3 sentences. Focus on the main topic, purpose, and key points. Maximum 200 words."
                            }, {
                                "role": "user",
                                "content": f"Summarize this document:\n\nFilename: {filename}\n\nContent:\n{text_for_summary}"
                            }],
                            "model": settings.OLLAMA_MODEL,
                            "max_tokens": 300,
                            "temperature": 0.3
                        },
                        headers={
                            "X-API-Key": settings.MICROSERVICES_API_KEY,
                            "Content-Type": "application/json"
                        },
                        timeout=20.0
                    )
                    
                    if response.status_code == 200:
                        result = response.json()
                        summary = result.get("choices", [{}])[0].get("message", {}).get("content", "")
                        if summary:
                            logger.info(f"Generated summary for {filename}: {len(summary)} chars")
                            return summary
                    
            except Exception as e:
                logger.warning(f"Failed to generate LLM summary: {e}")
            
            # Fallback to simple extraction if LLM fails
            return self._create_simple_summary(text, filename)
            
        except Exception as e:
            logger.error(f"Error generating summary: {e}")
            return self._create_simple_summary(text, filename)
    
    def _create_simple_summary(self, text: str, filename: str) -> str:
        """Create a simple summary without LLM"""
        # Clean up text
        lines = [line.strip() for line in text.split('\n') if line.strip()]
        
        # Use DocumentTypeDetector for accurate type detection
        try:
            from app.services.document_type_detector import get_document_type_detector
            detector = get_document_type_detector()
            doc_type, confidence = detector.detect_type(text, filename)
            
            # Capitalize first letter for display
            doc_type_display = doc_type.capitalize()
            
            # Add confidence indicator if low
            if confidence < 0.5:
                doc_type_display = f"Possible {doc_type_display}"
        except Exception as e:
            logger.warning(f"Could not use DocumentTypeDetector: {e}")
            # Fallback to basic detection
            doc_type_display = "Document"
        
        # Get first meaningful lines
        meaningful_lines = []
        for line in lines[:10]:
            if len(line) > 20:  # Skip very short lines
                meaningful_lines.append(line)
                if len(meaningful_lines) >= 3:
                    break
        
        if meaningful_lines:
            preview = " ".join(meaningful_lines[:2])[:200]
            return f"{doc_type_display}: {filename}. {preview}..."
        else:
            return f"{doc_type_display}: {filename}. Content preview: {text[:200]}..."
    
    async def _perform_routing_analysis(self, db: AsyncSession, doc_info: dict, text: str, file_ext: str):
        """Perform routing analysis using the same DB session"""
        try:
            from app.services.agent_router_service import AgentRouterService
            
            # Initialize router service
            router_service = AgentRouterService(
                tenant_id=doc_info["tenant_id"],
                user_id=doc_info["user_id"]
            )
            
            # Analyze and route document
            routing_result = await router_service.analyze_and_route_document(
                db=db,
                document_id=doc_info["id"],
                content=text,
                filename=doc_info["filename"],
                file_type=file_ext
            )
            
            if routing_result.get("success"):
                logger.info(
                    f"Document {doc_info['id']} routed successfully: "
                    f"Type: {routing_result.get('document_type')}, "
                    f"Agents: {len(routing_result.get('assigned_agents', []))}"
                )
            else:
                logger.warning(f"Document routing failed for {doc_info['id']}: {routing_result.get('error')}")
                
        except Exception as e:
            logger.warning(f"Failed to route document {doc_info['id']} with Agent Router: {e}")
            # Don't fail the whole process if routing fails